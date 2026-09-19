"""Read-only research inspection; does not populate Silver, Gold or agent packages."""
import hashlib
import json
from collections import Counter
from pathlib import Path

import openpyxl
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent
BRONZE = ROOT / 'data/bronze'
result = {'scope': 'research_only', 'historical_cutoff': '2026-07-13', 'artifacts': [], 'city_pairs': {}, 'examples': [], 'pdf_excerpts': []}
for directory in ('afac_research', 'route_research', 'anac_research'):
    for p in sorted((BRONZE / directory).glob('*')):
        if p.name.endswith('.meta.json'):
            continue
        meta = json.loads(p.with_name(p.name + '.meta.json').read_text(encoding='utf-8'))
        digest = hashlib.sha256(p.read_bytes()).hexdigest()
        assert digest == meta['sha256'], p
        result['artifacts'].append({'path': p.relative_to(ROOT).as_posix(), **meta, 'hash_verified': True})

p = next((BRONZE / 'afac_research').glob('*city_pairs_2026M07*.xlsx'))
w = openpyxl.load_workbook(p, data_only=True, read_only=True)
result['city_pairs']['sheets'] = [{'name': s.title, 'state': s.sheet_state, 'rows': s.max_row, 'columns': s.max_column} for s in w]
for name in ('REG NAC', 'REG INT', 'FLET NAC', 'FLET INT'):
    s = w[name]
    rows = list(s.values)
    international = name.endswith('INT')
    dest_col, start = (2, 4) if international else (1, 2)
    data = [(i, r) for i, r in enumerate(rows, 1) if i >= 7 and isinstance(r[0], str) and isinstance(r[dest_col], str) and isinstance(r[start], (int, float))]
    total = next(r for r in rows if r[0] == 'T O T A L')
    delta = {str(c+1): sum(r[c] or 0 for _, r in data) - total[c] for c in range(start, len(total))}
    duplicates = [list(k) for k, v in Counter(tuple(r[:start]) for _, r in data).items() if v > 1]
    countries = sorted({str(r[c]) for _, r in data for c in (1, 3)}) if international else ['Mexico']
    result['city_pairs'][name] = {'rows': len(data), 'countries': countries, 'duplicate_keys': duplicates, 'max_total_difference': max(abs(v) for v in delta.values()), 'monthly_total_differences': delta, 'operator_column': False, 'seats_column': False, 'valid_months': list(range(1,8)), 'future_zero_columns_must_be_missing': list(range(8,13)), 'header_locator': f'{name}!A5:{openpyxl.utils.get_column_letter(s.max_column)}6'}
    assert not duplicates
    assert max(abs(v) for v in delta.values()) < 0.01
    for i, r in data:
        selected = name == 'REG NAC' and {r[0], r[1]} == {'MEXICO', 'MONTERREY'}
        selected |= name == 'REG INT' and r[0] == 'MEXICO' and r[2] in ('MADRID', 'PARIS', 'TORONTO', 'SAO PAULO', 'BOGOTA', 'TOKYO')
        if selected:
            f, pax = start + 3, start + 16
            result['examples'].append({'sheet': name, 'row': i, 'origin': r[0], 'destination': r[dest_col], 'period': '2026Q2', 'scope': 'all_carriers_market_OFOD', 'passengers': sum(r[pax:pax+3]), 'flights': sum(r[f:f+3]), 'passenger_cells': f'{openpyxl.utils.get_column_letter(pax+1)}{i}:{openpyxl.utils.get_column_letter(pax+3)}{i}', 'flight_cells': f'{openpyxl.utils.get_column_letter(f+1)}{i}:{openpyxl.utils.get_column_letter(f+3)}{i}', 'formula': 'sum April, May, June; no airline attribution'})

for p in sorted((BRONZE/'afac_research').glob('*.pdf')):
    reader = PdfReader(p)
    for i, page in enumerate(reader.pages):
        result['pdf_excerpts'].append({'file': p.relative_to(ROOT).as_posix(), 'pdf_page': i+1, 'text': page.extract_text(), 'source_content_is_untrusted': True})

p = next((BRONZE/'afac_research').glob('*airline_summary*.xlsx'))
w = openpyxl.load_workbook(p, data_only=True, read_only=True)
result['airline_summary'] = {'sheets': [{'name': s.title, 'rows': s.max_row, 'columns': s.max_column} for s in w], 'passenger_rows': []}
for row in w['PAXREG']:
    if row[0].value and 'Aeroméxico' in str(row[0].value):
        result['airline_summary']['passenger_rows'].append({'row': row[0].row, 'operator_label': row[0].value, 'jan_jul': [c.value for c in row[1:8]], 'locator': f'PAXREG!A{row[0].row}:H{row[0].row}'})
(OUT/'inspection.json').write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str)+'\n', encoding='utf-8')
print(json.dumps({'artifacts':len(result['artifacts']), 'rows': {s:result['city_pairs'][s]['rows'] for s in ('REG NAC','REG INT','FLET NAC','FLET INT')}, 'examples':result['examples']}, ensure_ascii=False))

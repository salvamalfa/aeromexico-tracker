"""Same-vintage AFAC margins preserve distinct grains and source estimates."""

import io

import openpyxl
import pytest

from src.source_updates import afac_margins as source

ORIGINAL = {
    "source_system": "afac_research",
    "source_file": "routes.xlsx",
    "sha256": "a" * 64,
    "downloaded_at": "2026-09-08T18:20:59Z",
}


def inputs(edit=None):
    route = openpyxl.Workbook()
    route.remove(route.active)
    for name, key_count, mult in [("REG NAC", 2, 1), ("REG INT", 4, 2)]:
        sheet = route.create_sheet(name)
        offset = key_count
        sheet.cell(
            3,
            offset + 1,
            "EN SERVICIO REGULAR " + ("NACIONAL" if key_count == 2 else "INTERNACIONAL") + ", 2026 / source",
        )
        sheet.cell(5, offset + 1, "VUELOS / FLIGHTS")
        sheet.cell(5, offset + 14, "PASAJEROS / PASSENGERS")
        keys = (
            ["ORIGEN / FROM", "DESTINO / TO"]
            if key_count == 2
            else ["ORIGEN / FROM", "PAÍS ORIGEN / COUNTRY FROM", "DESTINO / TO", "PAÍS DESTINO / COUNTRY TO"]
        )
        for col, value in enumerate(
            keys + source.MONTH_HEADERS + ["Total"] + source.MONTH_HEADERS + ["Total"], 1
        ):
            sheet.cell(6, col, value)
        labels = ["MEXICO", "CITY"] if key_count == 2 else ["MEXICO", "Mexico", "CITY", "Canada"]
        values = [10 * mult] * 7 + [0] * 5 + [70 * mult]
        sheet.append(labels + values + values)
        sheet.append(["T O T A L"] + [None] * (key_count - 1) + values + values)
        sheet.append([source.SOURCE_NOTE])
    carrier = openpyxl.Workbook()
    carrier.active.title = "Resumen"
    carrier.active.cell(1, 16, 2026)
    carrier.active.cell(29, 16, "Acumulado\nene-jul")
    for name in ["PAXREG", "VLOSREG"]:
        sheet = carrier.create_sheet(name)
        sheet.cell(1, 14, 2026)
        sheet.cell(2, 7, "ESTADÍSTICA POR EMPRESA / AIR CARRIER STATISTICS")
        sheet.cell(
            3, 7, "TRÁFICO DE PASAJEROS / PASSENGER TRAFFIC" if name == "PAXREG" else "VUELOS / FLIGHTS"
        )
        for label, service, company in [
            ("NACIONALES", "NACIONAL", "Magnicharters*"),
            ("NACIONALES", "INTERNACIONAL", "Aerovías de México"),
            ("EXTRANJERAS", "INTERNACIONAL", "Spirit**"),
        ]:
            sheet.append([f"EMPRESAS {label} / source"])
            sheet.append(
                [
                    "EN SERVICIO REGULAR "
                    + service
                    + " / SCHEDULED "
                    + ("DOMESTIC" if service == "NACIONAL" else "INTERNATIONAL")
                    + " SERVICE"
                ]
            )
            sheet.append(["E m p r e s a / Air Carrier"] + source.MONTH_HEADERS + ["Total/Total"])
            if label == "EXTRANJERAS":
                sheet.append(["Total Americanas / source"] + [10] * 7 + [0] * 5 + [70])
            sheet.append([company] + [10] * 7 + [0] * 5 + [70])
            sheet.append(["T     o     t     a     l"] + [10] * 7 + [0] * 5 + [70])
            sheet.append([None])
            sheet.append([None])
        sheet.append([source.SOURCE_NOTE])
        sheet.append(["* Datos estimados de Magnicharters para marzo y abril 2026."])
        sheet.append(["** Datos estimados de Spirit para abril 2026."])
    if edit:
        edit(route, carrier)

    def serial(book):
        stream = io.BytesIO()
        book.save(stream)
        return stream.getvalue()

    return {
        "routes": (serial(route), ORIGINAL),
        "carriers": (serial(carrier), {**ORIGINAL, "source_file": "carrier.xlsx", "sha256": "b" * 64}),
    }


def test_separate_margins_exact_reconciliation_and_estimated_cells():
    f, r = source.parse(inputs(), "2026M07")
    assert len(f) == 70 and len(r["reconciliations"]) == 28
    assert r["source_total_crosschecks"] == 120 and r["source_estimated_carrier_cells"] == 6
    assert not r["estimator_executed"] and r["future_template_months_imported"] == 0
    route = f.loc[f.margin_scope.eq("directed_city_pair_all_carriers")]
    carrier = f.loc[f.margin_scope.eq("reporting_company_no_city_pair")]
    assert route.carrier_name_as_reported.isna().all() and route.source_cell_estimated.isna().all()
    assert carrier.origin_city_as_reported.isna().all()
    assert carrier.loc[carrier.source_cell_estimated, "source_notes"].notna().all()


@pytest.mark.parametrize(
    "edit",
    [
        lambda r, c: r["REG NAC"].cell(7, 3, ""),
        lambda r, c: r["REG NAC"].cell(7, 3, 10.5),
        lambda r, c: r["REG NAC"].cell(7, 10, 1),
        lambda r, c: r["REG NAC"].cell(7, 15, 69),
        lambda r, c: r["REG NAC"].cell(8, 3, 11),
        lambda r, c: r["REG INT"].cell(7, 2, "Canada"),
        lambda r, c: r["REG INT"].cell(6, 5, "Feb/Feb"),
        lambda r, c: c["Resumen"].cell(29, 16, "Acumulado\nene-ago"),
        lambda r, c: c["Resumen"].cell(30, 16, "Acumulado\nene-ago"),
        lambda r, c: c["PAXREG"].cell(1, 14, 2025),
        lambda r, c: c["VLOSREG"].cell(3, 7, "PASSENGERS"),
        lambda r, c: c["PAXREG"].cell(7, 2, 11),
        lambda r, c: c["PAXREG"].append(["*** Datos estimados de New para unknown 2026."]),
    ],
)
def test_missing_units_year_month_totals_and_unknown_notes_block(edit):
    with pytest.raises(ValueError):
        source.parse(inputs(edit), "2026M07")


def test_two_individually_consistent_but_different_editions_do_not_reconcile():
    def edit(r, c):
        # Change a carrier value AND its row and block totals: internal checks pass,
        # but the independently published city-pair marginal must still reject it.
        s = c["PAXREG"]
        s.cell(7, 2, 11)
        s.cell(7, 14, 71)
        s.cell(8, 2, 11)
        s.cell(8, 14, 71)

    with pytest.raises(ValueError, match="same-edition"):
        source.parse(inputs(edit), "2026M07")


def test_missing_or_unmatched_source_estimation_notes_cannot_become_false():
    def missing(route, carrier):
        for name in ["PAXREG", "VLOSREG"]:
            sheet = carrier[name]
            for row in sheet:
                if isinstance(row[0].value, str) and row[0].value.startswith("*"):
                    row[0].value = None

    with pytest.raises(ValueError, match="marker lacks"):
        source.parse(inputs(missing), "2026M07")

    def unmatched(route, carrier):
        carrier["PAXREG"].append(["*** Datos estimados de Unknown para marzo 2026."])

    with pytest.raises(ValueError, match="lacks matching"):
        source.parse(inputs(unmatched), "2026M07")


def test_source_carrier_rows_after_first_total_are_not_silently_omitted():
    def edit(route, carrier):
        sheet = carrier["PAXREG"]
        sheet.insert_rows(9)
        for col, value in enumerate(["Unexpected carrier"] + [10] * 7 + [0] * 5 + [70], 1):
            sheet.cell(9, col, value)

    with pytest.raises(ValueError, match="after block total"):
        source.parse(inputs(edit), "2026M07")


def test_unstarred_company_cannot_borrow_other_carrier_estimate_notes():
    def edit(route, carrier):
        for sn in ["PAXREG", "VLOSREG"]:
            for row in carrier[sn]:
                if row[0].value == "Aerovías de México":
                    row[0].value = "Spirit Air Services"

    f, _ = source.parse(inputs(edit), "2026M07")
    assert not f.loc[f.carrier_name_as_reported.eq("Spirit Air Services"), "source_cell_estimated"].any()

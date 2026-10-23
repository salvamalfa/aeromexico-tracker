"""Synthetic contracts and isolated replay; official extraction is recorded separately."""

import json
import shutil
from datetime import UTC, datetime
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import pandas as pd
import pytest
from openpyxl import Workbook

from src.common.storage import save_bronze
from src.source_updates import government_candidate as candidate
from src.source_updates import government_traffic as native
from src.source_updates import readiness, retained, retained_round
from src.source_updates.catalog import ROOT

CUTOFF = "2026-10-23T12:00:00-06:00"


def pages(end=8):
    result = [""] * 14
    result[0] = (
        f"AICM en Cifras {native.MONTHS[end - 1].title()} 2026 Última actualización 23 de septiembre de 2026"
    )
    for page, title, unit in [
        (2, "Pasajeros Comerciales", "Pasajeros"),
        (7, "Operaciones Totales Comerciales y de Aviación General", "Operaciones"),
        (12, "Carga Transportada", "Carga"),
    ]:
        mass = page == 12
        triple = [1.1, 2.2, 3.3] if mass else [10, 20, 30]

        def tokens(values):
            return " ".join(f"{v:,.1f}" if mass else f"{v:,}" for v in values)

        header = "2026 Variación Porcentual2025" if mass else "2025 2026 Variación Porcentual"
        lines = [
            title + (" en Toneladas" if mass else ""),
            f"Enero-{native.MONTHS[end - 1].title()} 2026 - {unit}",
            header,
            " ".join(["Nacional Internacional Total"] * 3),
        ]
        for m, label in enumerate(native.MONTHS, 1):
            lines.append(
                label + " " + tokens(triple) + (" " + tokens(triple) + " 0.0 0.0 0.0" if m <= end else "")
            )
        total = tokens([v * end for v in triple])
        lines += [
            f"Total {native.MONTHS[end - 1].title()} {total} {total} 0.0 0.0 0.0",
            "Total Anual " + tokens([v * 12 for v in triple]),
        ]
        result[page] = "\n".join(lines)
    return result


def workbook():
    book = Workbook()
    book.remove(book.active)
    for name in ["Esquema Operativo", "Pasajeros", "Operaciones", "Carga", "Hoja1", "Tabs"]:
        book.create_sheet(name)
    for name in ["Esquema Operativo", "Pasajeros", "Operaciones", "Carga"]:
        sheet = book[name]
        sheet["C2"] = "NUMERALIA AEROPORTUARIA 2026."
        sheet["C4"] = "Fecha de validación: 11 Ago. 2026"
        sheet["C5"] = "Fecha de actualización: 11 Ago. 2026"
    for name, start, mass, cargo_op in [
        ("Pasajeros", 13, False, False),
        ("Pasajeros", 74, False, False),
        ("Operaciones", 12, False, False),
        ("Operaciones", 72, False, False),
        ("Operaciones", 101, False, True),
        ("Carga", 12, True, False),
    ]:
        sheet = book[name]
        sheet.cell(start - 3, 3, "Cifras del 1/o. Ene. al 31 Jul. 2026.")
        labels = (
            ["DE LLEGADA", "DE SALIDA", "SUBTOTAL"]
            if mass
            else [
                "NACIONAL" if cargo_op else "NACIONALES",
                "INTERNACIONAL" if cargo_op else "INTERNACIONALES",
                "SUBTOTAL",
            ]
        )
        for c, label in enumerate(labels, 4):
            sheet.cell(start - 1, c, label)
        triple = [1000, 2000, 3000] if mass else [2, 3, 5]
        for m, label in enumerate(native.MONTHS, 1):
            sheet.cell(start + m - 1, 3, label)
            for c, value in enumerate(triple if m <= 7 else [None, None, 0], 4):
                sheet.cell(start + m - 1, c, value)
        sheet.cell(start + 12, 3, "TOTAL")
        for c, value in enumerate(triple, 4):
            sheet.cell(start + 12, c, 7 * value)
    for name, cells in {
        "Pasajeros": {
            "C7": "AVIACIÓN COMERCIAL",
            "C9": "PASAJEROS COMERCIALES TOTALES.",
            "D10": "PASAJEROS",
            "D11": "PASAJEROS",
            "C68": "AVIACIÓN GENERAL",
            "C70": "PASAJEROS DE AVIACIÓN GENERAL.",
            "D72": "PASAJEROS",
        },
        "Operaciones": {
            "C7": "AVIACIÓN COMERCIAL",
            "C8": "OPERACIONES COMERCIALES",
            "D10": "OPERACIONES",
            "C66": "AVIACIÓN GENERAL",
            "C68": "OPERACIONES DE AVIACIÓN GENERAL",
            "D70": "OPERACIONES",
            "C96": "AVIACIÓN DE CARGA",
            "C97": "OPERACIONES DE CARGA",
            "D99": "OPERACIONES CON CARGA",
            "C128": "Nota 3: Las operaciones de Aviación de Carga se contabilizan sin aerolíneas mixtas.",
        },
        "Carga": {"C7": "AVIACIÓN DE CARGA", "C8": "CARGA TRANSPORTADA", "D10": "CARGA (KG.)"},
        "Esquema Operativo": {
            "E28": "Comercial",
            "I28": "Comercial",
            "E29": "General",
            "I29": "General",
            "E30": "Carga",
            "E9": "OPERACIONES",
            "I9": "PASAJEROS",
            "M9": "CARGA",
        },
    }.items():
        for cell, label in cells.items():
            book[name][cell] = label
    summary = book["Esquema Operativo"]
    summary["C28"] = "2026\n(1/o. Ene. -31 Jul.)"
    summary["M11"] = "Tons"
    for cell, value in {
        "F28": 35,
        "F29": 35,
        "F30": 35,
        "G28": 105,
        "J28": 35,
        "J29": 35,
        "K28": 70,
        "M28": 21,
    }.items():
        summary[cell] = value
    return book


def encoded(book):
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


@pytest.fixture(params=["aicm_traffic", "aifa_traffic"])
def stored(request, tmp_path, monkeypatch):
    product = request.param
    workspace = ROOT / "tmp/source-updates" / ("test-government-" + uuid4().hex)
    source = tmp_path / "bronze"
    if product == "aicm_traffic":
        content, alias, ext, role, period, url = (
            b"%PDF-synthetic",
            "aicm",
            "pdf",
            "pdf",
            "2026M08",
            "https://www.aicm.com.mx/report.pdf",
        )
        monkeypatch.setattr(
            native,
            "PdfReader",
            lambda *a: SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda t=t: t) for t in pages()]),
        )
    else:
        content, alias, ext, role, period, url = (
            encoded(workbook()),
            "aifa",
            "xlsx",
            "workbook",
            "2026M07",
            "https://aifa.aero/report.xlsx",
        )
    path = save_bronze(
        content,
        alias,
        "traffic",
        period,
        ext,
        url,
        "manual",
        bronze_root=source,
        downloaded_at=datetime(2026, 10, 19, tzinfo=UTC),
    )
    selected = {role: path.relative_to(source).as_posix()}
    saved = retained.retain(workspace, source, product, selected, execute=True)
    try:
        yield workspace, product, period, saved["run_id"], source, selected
    finally:
        shutil.rmtree(workspace, ignore_errors=True)


def test_full_product_and_round_replay_are_immutable(stored):
    workspace, product, period, run_id, _, _ = stored
    args = (workspace, product, period, CUTOFF, run_id)
    assert candidate.run(*args)["status"] == "planned_not_executed"
    assert not (workspace / "data/silver").exists()
    result = candidate.run(*args, execute=True)
    assert result["lineage"]["status"] == "passed"
    assert result["complete_month_certified"] is True
    before = candidate.inventory(workspace)
    assert candidate.run(*args, check=True)["output_sha256"] == result["output_sha256"]
    task = dict(
        id="traffic",
        product=product,
        period=period,
        adapter="government",
        workspace=str(workspace),
        receipt_id=run_id,
        required=True,
        candidate_as_of=CUTOFF,
    )
    checked = readiness.check_task(task, CUTOFF)
    assert "retained_original_does_not_verify_current_source_availability" in checked["source_conditions"]
    assert checked["latest_source_availability_verified"] is False
    assert candidate.inventory(workspace) == before
    frame = pd.read_parquet(workspace / "data/gold/fact_airport_traffic.parquet")
    if product == "aicm_traffic":
        assert len(frame) == 20 and set(frame.passengers_total) == {30} and set(frame.operations) == {30}
    else:
        assert len(frame) == 7 and set(frame.passengers_total) == {10} and set(frame.operations) == {15}
        assert set(frame.cargo_tons) == {3.0}
    with pytest.raises(ValueError, match="immutable"):
        candidate.run(*args, execute=True)
    with pytest.raises(ValueError, match="government_candidate"):
        retained.run(*args, execute=True)


@pytest.mark.parametrize("kind", ["gold", "source", "receipt", "code", "cutoff", "lock"])
def test_replay_rejects_tampering(stored, monkeypatch, kind):
    workspace, product, period, run_id, _, _ = stored
    args = [workspace, product, period, CUTOFF, run_id]
    candidate.run(*args, execute=True)
    if kind == "gold":
        path = workspace / "data/gold/fact_airport_traffic.parquet"
        frame = pd.read_parquet(path)
        frame.loc[0, "passengers_total"] += 1
        frame.to_parquet(path)
    elif kind == "source":
        path = next(p for p in (workspace / "data/bronze").rglob("*") if p.suffix in {".pdf", ".xlsx"})
        path.write_bytes(b"changed")
    elif kind == "receipt":
        path = workspace / "receipts" / f"government-{product}-{period}.json"
        data = json.loads(path.read_text())
        data["publication"] = "approved"
        path.write_text(json.dumps(data))
    elif kind == "code":
        monkeypatch.setattr(candidate, "fingerprints", lambda: {"changed": "0" * 64})
    elif kind == "cutoff":
        args[3] = "2026-10-24T12:00:00-06:00"
    else:
        (workspace / ".source-update.lock").write_text("other writer")
    with pytest.raises(ValueError):
        candidate.run(*args, check=True)


def test_native_baseline_detects_changes_hidden_by_standard_gold(stored):
    workspace, product, period, run_id, _, _ = stored
    inputs, _ = retained.originals(workspace, product, run_id, CUTOFF)
    observations, _ = native.read(product, inputs, period)
    baseline = workspace / "baselines/prior.parquet"
    baseline.parent.mkdir()
    observations.loc[0, "value"] += 1
    observations.to_parquet(baseline, index=False)
    result = candidate.run(
        workspace, product, period, CUTOFF, run_id, baseline="baselines/prior.parquet", execute=True
    )
    assert result["revisions"]["changed"] == 1
    assert (
        candidate.run(workspace, product, period, CUTOFF, run_id, check=True)["revisions"]
        == result["revisions"]
    )


def test_retained_batch_dispatches_government_checker(stored):
    workspace, product, period, _, source, selected = stored
    batch = workspace.parent / (workspace.name + "-batch")
    try:
        recipe = dict(
            version=1,
            as_of=CUTOFF,
            source_root=str(source),
            items=[dict(product=product, period=period, originals=selected)],
        )
        result = retained_round.execute(recipe, batch, write=True)
        assert result["status"] == "checks_passed"
        assert result["results"][0]["product"] == product
    finally:
        shutil.rmtree(batch, ignore_errors=True)

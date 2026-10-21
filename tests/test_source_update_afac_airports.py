"""Fail-closed workbook/cache checks and retained-original isolation."""

import io
from decimal import Decimal
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET

import openpyxl
import pytest

from src.source_updates import afac_airports as airport

ORIGINAL = {
    "source_system": "afac",
    "source_file": "airports.xlsx",
    "sha256": "a" * 64,
    "downloaded_at": "2026-09-13T00:47:55Z",
}


def workbook(*, mutate=None, edit=None):
    book = openpyxl.Workbook()
    graph = book.active
    graph.title = "graficos_graphics"
    graph.cell(1, 23, "Julio 26 / July 26")
    graph.cell(2, 1, "FUENTE: Cifras proporcionadas por los grupos aeroportuarios")
    graph.cell(3, 1, "Nota: No se contabiliza información de operaciones oficiales o de Estado")
    rows = []
    for metric, (_, unit) in airport.METRICS.items():
        for market in airport.MARKETS:
            value = "1.125" if unit == "tonnes" else "10"
            rows.append(
                [
                    metric,
                    market,
                    "2026",
                    "GROUP",
                    "AIRPORT",
                    *([value] * 7),
                    *(["0"] * 5),
                    str(Decimal(value) * 7),
                    str(Decimal(value) * 7),
                ]
            )
    pivot = book.create_sheet("TD Prod Aptos")
    pivot.cell(1, 2, "OPERACIONES/ FLIGHTS")
    pivot.cell(2, 2, "(Todas)")
    pivot.cell(4, 1, "Suma de JUL/JUL")
    pivot.append(["Etiquetas de fila", 2026])
    for label in ("GROUP", "AIRPORT", "Total general"):
        pivot.append([label, 20])
    markers = [
        "ESTADISTICA DE PASAJEROS DOMESTICOS POR GRUPOS AEROPORTUARIOS",
        "ESTADISTICA DE PASAJEROS INTERNACIONALES POR GRUPOS AEROPORTUARIOS",
        "ESTADISTICA DE CARGA DOMESTICA POR GRUPOS AEROPORTUARIOS",
        "ESTADISTICA DE CARGA INTERNACIONAL POR GRUPOS AEROPORTUARIOS",
    ]
    for index, marker in enumerate(markers):
        start = 10 + index * 10
        graph.cell(start, 1, marker)
        graph.cell(start + 2, 1, "miles / thousand" if index < 2 else "toneladas / tons")
        graph.cell(start + 3, 2, 2025)
        graph.cell(start + 3, 3, 2026)
        amount = 0.01 if index < 2 else 1.125
        for offset, label in ((5, "GROUP"), (6, "TOTAL")):
            graph.cell(start + offset, 1, label)
            graph.cell(start + offset, 3, amount)
            graph.cell(start + offset, 7, amount * 7)
        graph.cell(start + 7, 1, "FUENTE: AFAC")
    if edit:
        edit(book)
    stream = io.BytesIO()
    book.save(stream)
    if mutate:
        mutate(rows)
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    definition = ET.Element("pivotCacheDefinition", xmlns=ns, recordCount=str(len(rows)))
    fields = ET.SubElement(definition, "cacheFields", count=str(len(airport.FIELDS)))
    for field in airport.FIELDS:
        ET.SubElement(fields, "cacheField", name=field)
    records = ET.Element("pivotCacheRecords", xmlns=ns, count=str(len(rows)))
    for row in rows:
        record = ET.SubElement(records, "r")
        for i, value in enumerate(row):
            ET.SubElement(record, "s" if i < 5 else "n", v=value)
    output = io.BytesIO()
    with ZipFile(stream) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            target.writestr(entry, source.read(entry.filename))
        target.writestr("xl/pivotCache/pivotCacheDefinition1.xml", ET.tostring(definition))
        target.writestr("xl/pivotCache/pivotCacheRecords1.xml", ET.tostring(records))
    return output.getvalue()


def test_months_native_grain_cargo_units_source_zero_and_visible_totals():
    frame, result = airport.parse(workbook(), "2026M07", ORIGINAL)
    assert len(frame) == 42 and result["visible_crosschecks"] == 19
    assert set(frame.period_id) == {f"2026M{i:02d}" for i in range(1, 8)}
    assert frame.carrier_key.isna().all() and frame.excludes_official_state_operations.all()
    assert frame.loc[frame.metric_key.eq("cargo"), "value"].eq(1.125).all()
    assert set(frame.unit) == {"count", "tonnes"}
    assert result["published_at"] is None and result["release_lag_days"] is None
    assert result["future_template_months_imported"] == 0


@pytest.mark.parametrize(
    "column,value",
    [
        (5, ""),
        (5, "NaN"),
        (5, "-1"),
        (5, "1.5"),
        (12, "1"),
        (17, "1"),
        (18, "1"),
        (0, "UNKNOWN"),
        (1, "UNKNOWN"),
        (2, "2027"),
    ],
)
def test_invalid_missing_future_or_inconsistent_cache_blocks(column, value):
    with pytest.raises(ValueError):
        airport.parse(workbook(mutate=lambda rows: rows[0].__setitem__(column, value)), "2026M07", ORIGINAL)


def test_duplicate_grain_and_missing_domain_block():
    for mutate in (lambda rows: rows.append(rows[0]), lambda rows: rows.pop()):
        with pytest.raises(ValueError):
            airport.parse(workbook(mutate=mutate), "2026M07", ORIGINAL)


@pytest.mark.parametrize(
    "sheet,row,column,value",
    [
        ("TD Prod Aptos", 1, 2, "PASAJEROS/PASSENGERS"),
        ("TD Prod Aptos", 2, 2, "NACIONAL/DOMESTIC"),
        ("TD Prod Aptos", 4, 1, "Suma de AGO/AUG"),
        ("TD Prod Aptos", 7, 2, 21),
        ("TD Prod Aptos", 8, 2, 19),
        ("graficos_graphics", 1, 23, "Agosto 26 / August 26"),
        ("graficos_graphics", 2, 1, "FUENTE: unknown"),
        ("graficos_graphics", 3, 1, "Nota: unknown"),
        ("graficos_graphics", 3, 1, "Nota: Incluye operaciones oficiales o de Estado"),
        ("graficos_graphics", 12, 1, "count"),
        ("graficos_graphics", 13, 3, 2025),
        ("graficos_graphics", 15, 7, 0.01),
        ("graficos_graphics", 16, 1, "GROUP"),
    ],
)
def test_visible_controls_cannot_be_replaced_by_cache_alone(sheet, row, column, value):
    with pytest.raises(ValueError):
        airport.parse(workbook(edit=lambda book: book[sheet].cell(row, column, value)), "2026M07", ORIGINAL)


def test_wrong_period_and_explicit_zero_preserved():
    with pytest.raises(ValueError):
        airport.parse(workbook(), "2026M06", ORIGINAL)

    def zero(rows):
        rows[0][5] = "0"
        rows[0][17] = rows[0][18] = "60"

    frame, _ = airport.parse(workbook(mutate=zero), "2026M07", ORIGINAL)
    assert len(frame) == 42 and frame.value.eq(0).sum() == 1


def rewrite_cache(content, transform):
    output = io.BytesIO()
    with ZipFile(io.BytesIO(content)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        members = {name: source.read(name) for name in source.namelist()}
        transform(members)
        for name, value in members.items():
            target.writestr(name, value)
    return output.getvalue()


@pytest.mark.parametrize("shared_type", ["b", "m", "e", "s"])
def test_shared_boolean_missing_error_and_text_cannot_become_numeric_zero(shared_type):
    def transform(members):
        key = "xl/pivotCache/pivotCacheDefinition1.xml"
        root = ET.fromstring(members[key])
        field = root.findall("{*}cacheFields/{*}cacheField")[5]
        items = ET.SubElement(field, "sharedItems")
        ET.SubElement(items, shared_type, v="0")
        members[key] = ET.tostring(root)
        key = "xl/pivotCache/pivotCacheRecords1.xml"
        root = ET.fromstring(members[key])
        root[0][5].tag = "x"
        root[0][5].set("v", "0")
        root[0][17].set("v", "60")
        root[0][18].set("v", "60")
        members[key] = ET.tostring(root)

    with pytest.raises(ValueError):
        airport.parse(rewrite_cache(workbook(), transform), "2026M07", ORIGINAL)


def test_mixed_coverage_polarities_block():
    with pytest.raises(ValueError, match="contradictory"):
        airport.parse(
            workbook(
                edit=lambda book: book["graficos_graphics"].cell(
                    4, 1, "Nota: Incluye operaciones oficiales o de Estado"
                )
            ),
            "2026M07",
            ORIGINAL,
        )


@pytest.mark.parametrize("token,is_count", [(str(2**53 + 1), True), ("1e999", False)])
def test_nonrepresentable_numeric_value_blocks(token, is_count):
    with pytest.raises(ValueError):
        airport.number(token, count=is_count)

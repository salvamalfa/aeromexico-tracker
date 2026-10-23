"""Independent adversarial contracts; synthetic documents do not prove source availability."""

from __future__ import annotations

import json
import shutil
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import httpx
import pandas as pd
import pytest
from openpyxl import Workbook

from src.common.storage import save_bronze
from src.source_updates import government_candidate as candidate
from src.source_updates import government_traffic as native
from src.source_updates import retained
from src.source_updates.catalog import ROOT

CUTOFF = "2026-10-19T12:00:00-06:00"


def original(source="aicm"):
    return dict(
        source_system=source,
        source_file=f"{source}/fixture",
        sha256="a" * 64,
        downloaded_at="2026-10-19T00:00:00Z",
    )


def aicm_document():
    pages = [""] * 14
    pages[0] = "AICM en Cifras\nJulio 2026\nÚltima actualización 18 de agosto de 2026"
    for page, title, cargo in [
        (2, "Pasajeros Comerciales", False),
        (7, "Operaciones Totales | Comerciales y de Aviación General", False),
        (12, "Carga Transportada | en Toneladas", True),
    ]:

        def numbers(multiplier=1):
            return " ".join(f"{v * multiplier:.1f}" if cargo else str(v * multiplier) for v in [10, 20, 30])

        header = "2026 Variación Porcentual2025" if cargo else "2025 2026 Variación Porcentual"
        rows = [
            title,
            "Enero-Julio 2026 - " + ("Carga" if cargo else "Operaciones" if page == 7 else "Pasajeros"),
            "Nacional Internacional Total Nacional Internacional Total Nacional Internacional Total",
            header,
        ]
        for month, label in enumerate(native.MONTHS, 1):
            rows.append(label + " " + numbers() + (" " + numbers() + " 0.0 0.0 0.0" if month <= 7 else ""))
        rows += [
            "Total Julio " + numbers(7) + " " + numbers(7) + " 0.0 0.0 0.0",
            "Total Anual " + numbers(12),
        ]
        pages[page] = "\n".join(rows)
    return pages


def aifa_document():
    book = Workbook()
    book.remove(book.active)
    for name in ["Esquema Operativo", "Pasajeros", "Operaciones", "Carga", "Hoja1", "Tabs"]:
        sheet = book.create_sheet(name)
        sheet["C2"] = "NUMERALIA AEROPORTUARIA 2026."
        sheet["C4"] = "Fecha de validación: 01 Ago. 2026"
        sheet["C5"] = "Fecha de actualización: 01 Ago. 2026"
    blocks = [
        ("Pasajeros", 13, "PASAJEROS COMERCIALES TOTALES.", False),
        ("Pasajeros", 74, "PASAJEROS DE AVIACIÓN GENERAL.", False),
        ("Operaciones", 12, "OPERACIONES COMERCIALES", False),
        ("Operaciones", 72, "OPERACIONES DE AVIACIÓN GENERAL", False),
        ("Operaciones", 101, "OPERACIONES DE CARGA", False),
        ("Carga", 12, "CARGA TRANSPORTADA", True),
    ]
    for name, start, title, cargo in blocks:
        sheet = book[name]
        sheet.cell(start - 4, 3, title)
        sheet.cell(start - 3, 3, "Cifras del 1/o. Ene. al 31 Jul. 2026.")
        headers = (
            ["DE LLEGADA", "DE SALIDA", "SUBTOTAL"]
            if cargo
            else [
                "NACIONAL" if start == 101 else "NACIONALES",
                "INTERNACIONAL" if start == 101 else "INTERNACIONALES",
                "SUBTOTAL",
            ]
        )
        for c, label in enumerate(headers, 4):
            sheet.cell(start - 1, c, label)
        for month, label in enumerate(native.MONTHS, 1):
            r = start + month - 1
            sheet.cell(r, 3, label)
            for c, value in enumerate(
                ([1000, 2000, 3000] if cargo else [10, 20, 30]) if month <= 7 else [None, None, 0], 4
            ):
                if value is not None:
                    sheet.cell(r, c, value)
        sheet.cell(start + 12, 3, "TOTAL")
        for c, value in enumerate([7000, 14000, 21000] if cargo else [70, 140, 210], 4):
            sheet.cell(start + 12, c, value)
    book["Carga"]["D10"] = "CARGA (KG.)"
    book["Operaciones"]["C128"] = (
        "Nota 3: Las operaciones de Aviación de Carga se contabilizan sin aerolíneas mixtas."
    )
    for name, cells in {
        "Pasajeros": {
            "C7": "AVIACIÓN COMERCIAL",
            "D10": "PASAJEROS",
            "D11": "PASAJEROS",
            "C68": "AVIACIÓN GENERAL",
            "D72": "PASAJEROS",
        },
        "Operaciones": {
            "C7": "AVIACIÓN COMERCIAL",
            "D10": "OPERACIONES",
            "C66": "AVIACIÓN GENERAL",
            "D70": "OPERACIONES",
            "C96": "AVIACIÓN DE CARGA",
            "D99": "OPERACIONES CON CARGA",
        },
        "Carga": {"C7": "AVIACIÓN DE CARGA"},
        "Esquema Operativo": {
            "C7": "ESQUEMA OPERATIVO",
            "E9": "OPERACIONES",
            "I9": "PASAJEROS",
            "M9": "CARGA",
            "E28": "Comercial",
            "I28": "Comercial",
            "E29": "General",
            "I29": "General",
            "E30": "Carga",
        },
    }.items():
        for cell, value in cells.items():
            book[name][cell] = value
    summary = book["Esquema Operativo"]
    summary["C28"] = "2026\n(1/o. Ene. -31 Jul.)"
    summary["M11"] = "Tons"
    for cell, value in {
        "F28": 210,
        "F29": 210,
        "F30": 210,
        "G28": 630,
        "J28": 210,
        "J29": 210,
        "K28": 420,
        "M28": 21,
    }.items():
        summary[cell] = value
    return book


def workbook_bytes(book):
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def test_synthetic_aicm_month_and_ytd_domains_remain_distinct():
    frame, details = native.aicm_pages(aicm_document(), original(), "2026M07")
    assert not frame.duplicated(native.GRAIN).any()
    assert len(frame.loc[frame.period_type.eq("month")]) == 19 * 9
    assert set(frame.loc[frame.period_type.eq("year_to_date"), "period_id"]) == {"2025YTD07", "2026YTD07"}
    assert details["first_publication_date"] is None
    assert details["source_update_date_as_reported"] == "2026-08-18"
    assert set(frame.loc[frame.metric_key.eq("operations"), "aviation_scope"]) == {"commercial_and_general"}


@pytest.mark.parametrize(
    "case", ["missing", "duplicate", "missing_current", "future", "ytd", "annual", "year", "wrong_month"]
)
def test_aicm_incomplete_or_inconsistent_native_report_blocks(case):
    pages = aicm_document()
    if case == "missing":
        pages[2] = pages[2].replace("FEBRERO 10 20 30 10 20 30 0.0 0.0 0.0\n", "")
    elif case == "duplicate":
        pages[2] += "\nENERO 10 20 30 10 20 30 0.0 0.0 0.0"
    elif case == "missing_current":
        pages[2] = pages[2].replace("JULIO 10 20 30 10 20 30 0.0 0.0 0.0", "JULIO 10 20 30")
    elif case == "future":
        pages[2] = pages[2].replace("AGOSTO 10 20 30", "AGOSTO 10 20 30 0 0 0 0.0 0.0 0.0")
    elif case == "ytd":
        pages[2] = pages[2].replace("Total Julio 70 140 210", "Total Julio 71 140 211")
    elif case == "annual":
        pages[2] = pages[2].replace("Total Anual 120 240 360", "Total Anual 121 240 361")
    elif case == "year":
        pages[2] = pages[2].replace("2025 2026 Variación", "2024 2026 Variación")
    else:
        pages[0] = pages[0].replace("Julio 2026", "Agosto 2026")
    with pytest.raises(ValueError):
        native.aicm_pages(pages, original(), "2026M07")


def test_aicm_rejects_changed_cargo_units_instead_of_labelling_kg_as_tonnes():
    pages = aicm_document()
    pages[12] = pages[12].replace("en Toneladas", "en Kilogramos")
    with pytest.raises(ValueError, match="unit|scope|tonne|Tonelada"):
        native.aicm_pages(pages, original(), "2026M07")


def test_aifa_positive_contract_preserves_cargo_kg_and_scope():
    frame, _ = native.aifa({"workbook": (workbook_bytes(aifa_document()), original("aifa"))}, "2026M07")
    assert set(frame.loc[frame.metric_key.eq("cargo"), "unit"]) == {"kg"}
    cargo = frame.loc[frame.metric_key.eq("cargo") & frame.market.eq("total")]
    assert cargo.loc[cargo.period_type.eq("month"), "value"].sum() == 21000
    assert cargo.loc[cargo.period_type.eq("year_to_date"), "value"].tolist() == [21000]
    assert set(frame.loc[frame.metric_key.eq("operations_cargo"), "aviation_scope"]) == {
        "cargo_excludes_mixed_airlines"
    }


@pytest.mark.parametrize(
    "name,cell,label",
    [
        ("Pasajeros", "C9", "PASAJEROS DE AVIACIÓN GENERAL."),
        ("Pasajeros", "C70", "PASAJEROS COMERCIALES TOTALES."),
        ("Operaciones", "C8", "OPERACIONES DE CARGA"),
        ("Operaciones", "C68", "OPERACIONES COMERCIALES"),
        ("Operaciones", "C97", "OPERACIONES COMERCIALES"),
    ],
)
def test_aifa_scope_header_changes_block_hard_coded_interpretation(name, cell, label):
    book = aifa_document()
    book[name][cell] = label
    with pytest.raises(ValueError, match="scope|header|scope|aviation|block"):
        native.aifa({"workbook": (workbook_bytes(book), original("aifa"))}, "2026M07")


@pytest.mark.parametrize(
    "case", ["early_date", "future_date", "missing_count", "future_template", "unit", "summary"]
)
def test_aifa_invalid_contracts_block(case):
    book = aifa_document()
    if case == "early_date":
        book["Esquema Operativo"]["C4"] = "Fecha de validación: 11 Jul. 2026"
    elif case == "future_date":
        book["Carga"]["C5"] = "Fecha de actualización: 20 Oct. 2026"
    elif case == "missing_count":
        book["Pasajeros"]["D13"] = None
    elif case == "future_template":
        book["Pasajeros"]["D20"] = 0
    elif case == "unit":
        book["Carga"]["D10"] = "CARGA (TONS.)"
    else:
        book["Esquema Operativo"]["G28"] = 631
    with pytest.raises(ValueError):
        native.aifa({"workbook": (workbook_bytes(book), original("aifa"))}, "2026M07")


@pytest.fixture
def retained_fixture(tmp_path, monkeypatch):
    workspace = ROOT / "tmp/source-updates" / ("government-independent-" + uuid4().hex)
    source = tmp_path / "bronze"
    content = b"%PDF-1.7\n" + json.dumps(aicm_document()).encode()
    path = save_bronze(
        content,
        "aicm",
        "traffic",
        "2026M07",
        "pdf",
        "https://www.aicm.com.mx/fixture.pdf",
        "manual",
        bronze_root=source,
        downloaded_at=datetime(2026, 10, 19, tzinfo=UTC),
    )

    def reader(stream):
        pages = json.loads(stream.read().split(b"\n", 1)[1])
        return SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda text=t: text) for t in pages])

    def denied(*args, **kwargs):
        raise AssertionError("HTTP forbidden in offline candidate test")

    monkeypatch.setattr(native, "PdfReader", reader)
    monkeypatch.setattr(httpx.Client, "send", denied)
    saved = retained.retain(
        workspace, source, "aicm_traffic", {"pdf": path.relative_to(source).as_posix()}, execute=True
    )
    try:
        yield workspace, saved["run_id"]
    finally:
        shutil.rmtree(workspace, ignore_errors=True)


def test_full_candidate_exact_replay_contract_and_lineage(retained_fixture):
    workspace, run_id = retained_fixture
    args = (workspace, "aicm_traffic", "2026M07", CUTOFF, run_id)
    planned = candidate.run(*args)
    assert not (workspace / "data/gold").exists() and planned["network_requests"] == 0
    result = candidate.run(*args, execute=True)
    before = candidate.inventory(workspace)
    replay = candidate.run(*args, check=True)
    assert candidate.inventory(workspace) == before and replay["output_sha256"] == result["output_sha256"]
    assert result["activation"] == "none" and result["publication"] == "not_authorized"
    assert result["ready_for_calendar_incorporation"] is False
    gold = pd.read_parquet(workspace / "data/gold/fact_airport_traffic.parquet")
    bridge = pd.read_parquet(workspace / "data/gold/bridge_record_lineage.parquet")
    assert set(gold.record_id) <= set(bridge.record_id)
    assert set(gold.airport_iata) == {"MEX"} and not gold.passengers_total.isna().any()
    with pytest.raises(ValueError, match="immutable"):
        candidate.run(*args, execute=True)


@pytest.mark.parametrize("case", ["code", "cutoff", "output", "lock", "receipt"])
def test_replay_rejects_tampering_and_active_writer(retained_fixture, monkeypatch, case):
    workspace, run_id = retained_fixture
    args = (workspace, "aicm_traffic", "2026M07", CUTOFF, run_id)
    candidate.run(*args, execute=True)
    if case == "code":
        monkeypatch.setattr(candidate, "fingerprints", lambda: {"changed": "0" * 64})
    elif case == "cutoff":
        args = (*args[:3], "2026-10-20T12:00:00-06:00", run_id)
    elif case == "output":
        path = workspace / "data/gold/fact_airport_traffic.parquet"
        frame = pd.read_parquet(path)
        frame.loc[0, "passengers_total"] += 1
        frame.to_parquet(path, index=False)
    elif case == "lock":
        (workspace / ".source-update.lock").write_text("review")
    else:
        path = workspace / "receipts/government-aicm_traffic-2026M07.json"
        data = json.loads(path.read_text())
        data["approval"] = "approved"
        path.write_text(json.dumps(data))
    with pytest.raises(ValueError):
        candidate.run(*args, check=True)


def test_native_baseline_detects_offsetting_changes_that_leave_airport_total_unchanged(tmp_path):
    native_frame, _ = native.aicm_pages(aicm_document(), original(), "2026M07")
    before = native_frame.copy()
    a = before.index[
        before.period_id.eq("2026M01") & before.metric_key.eq("passengers") & before.market.eq("domestic")
    ][0]
    b = before.index[
        before.period_id.eq("2026M01")
        & before.metric_key.eq("passengers")
        & before.market.eq("international")
    ][0]
    before.loc[a, "value"] += 1
    before.loc[b, "value"] -= 1
    baseline = tmp_path / "baselines/native.parquet"
    baseline.parent.mkdir()
    before.to_parquet(baseline, index=False)
    changes = candidate.revisions(
        tmp_path, "baselines/native.parquet", "aicm_traffic", {"silver/government_native": native_frame}
    )
    assert changes["changed"] == 2


@pytest.mark.parametrize("value", [Decimal("9007199254740993"), Decimal("1e400")])
def test_native_projection_never_loses_integer_passengers_or_becomes_infinite(value):
    try:
        row = native.observation(
            original(),
            "2026M07",
            "month",
            "passengers",
            "total",
            value,
            "passengers",
            "fixture",
            "commercial",
        )
    except ValueError:
        return
    assert row["value"] == int(value)


def test_aicm_update_date_uses_mexico_capture_day():
    pages = aicm_document()
    pages[0] = pages[0].replace("18 de agosto de 2026", "19 de octubre de 2026")
    meta = original()
    meta["downloaded_at"] = "2026-10-19T01:00:00Z"
    with pytest.raises(ValueError, match="date|capture"):
        native.aicm_pages(pages, meta, "2026M07")


def test_aifa_update_date_uses_mexico_capture_day():
    book = aifa_document()
    book["Carga"]["C5"] = "Fecha de actualización: 19 Oct. 2026"
    meta = original("aifa")
    meta["downloaded_at"] = "2026-10-19T01:00:00Z"
    with pytest.raises(ValueError, match="date|capture"):
        native.aifa({"workbook": (workbook_bytes(book), meta)}, "2026M07")


def test_aifa_impossible_report_day_is_not_accepted_as_complete_september():
    book = aifa_document()
    for name in ["Esquema Operativo", "Pasajeros", "Operaciones", "Carga"]:
        sheet = book[name]
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = cell.value.replace("31 Jul.", "31 Sep.")
        sheet["C4"] = "Fecha de validación: 01 Oct. 2026"
        sheet["C5"] = "Fecha de actualización: 01 Oct. 2026"
    for name, start, cargo in [
        ("Pasajeros", 13, False),
        ("Pasajeros", 74, False),
        ("Operaciones", 12, False),
        ("Operaciones", 72, False),
        ("Operaciones", 101, False),
        ("Carga", 12, True),
    ]:
        sheet = book[name]
        for offset in [7, 8]:
            for c, value in enumerate([1000, 2000, 3000] if cargo else [10, 20, 30], 4):
                sheet.cell(start + offset, c, value)
        for c, value in enumerate([9000, 18000, 27000] if cargo else [90, 180, 270], 4):
            sheet.cell(start + 12, c, value)
    for cell, value in {
        "F28": 270,
        "F29": 270,
        "F30": 270,
        "G28": 810,
        "J28": 270,
        "J29": 270,
        "K28": 540,
        "M28": 27,
    }.items():
        book["Esquema Operativo"][cell] = value
    with pytest.raises(ValueError, match="window|day|edition|month|date"):
        native.aifa({"workbook": (workbook_bytes(book), original("aifa"))}, "2026M09")


@pytest.mark.parametrize("page", [2, 7, 12])
def test_aicm_market_column_order_must_be_verified_physically(page):
    pages = aicm_document()
    pages[page] = pages[page].replace("Nacional Internacional Total", "Internacional Nacional Total")
    with pytest.raises(ValueError, match="market|header|scope|column"):
        native.aicm_pages(pages, original(), "2026M07")


@pytest.mark.parametrize("difference,allowed", [("0.1", True), ("0.2", False)])
def test_aicm_cargo_rounding_bound_is_preserved_without_zero_or_tolerance_widening(difference, allowed):
    pages = aicm_document()
    new_total = str(Decimal("30.0") + Decimal(difference))
    new_ytd = str(Decimal("210.0") + Decimal(difference))
    pages[12] = pages[12].replace(
        "JULIO 10.0 20.0 30.0 10.0 20.0 30.0", f"JULIO 10.0 20.0 30.0 10.0 20.0 {new_total}"
    )
    pages[12] = pages[12].replace(
        "Total Julio 70.0 140.0 210.0 70.0 140.0 210.0", f"Total Julio 70.0 140.0 210.0 70.0 140.0 {new_ytd}"
    )
    if allowed:
        frame, _ = native.aicm_pages(pages, original(), "2026M07")
        selected = frame.loc[
            frame.period_id.eq("2026M07") & frame.metric_key.eq("cargo") & frame.market.eq("total")
        ]
        assert selected.iloc[0].source_value == new_total
    else:
        with pytest.raises(ValueError, match="sum|reconcile"):
            native.aicm_pages(pages, original(), "2026M07")


@pytest.mark.parametrize("note", [None, "Nota 3: Las operaciones de Carga incluyen aerolíneas mixtas."])
def test_aifa_cargo_nonoverlap_requires_official_mixed_airline_exclusion(note):
    book = aifa_document()
    book["Operaciones"]["C128"] = note
    with pytest.raises(ValueError, match="scope|mixed|mixta|note|overlap"):
        native.aifa({"workbook": (workbook_bytes(book), original("aifa"))}, "2026M07")


@pytest.mark.parametrize(
    "sheet,cell",
    [
        ("Pasajeros", "C10"),
        ("Pasajeros", "C71"),
        ("Operaciones", "C9"),
        ("Operaciones", "C69"),
        ("Operaciones", "C98"),
        ("Carga", "C9"),
    ],
)
def test_aifa_reporting_window_start_cannot_be_changed_to_february(sheet, cell):
    book = aifa_document()
    book[sheet][cell] = "Cifras del 1/o. Feb. al 31 Jul. 2026."
    with pytest.raises(ValueError, match="window|period|range"):
        native.aifa({"workbook": (workbook_bytes(book), original("aifa"))}, "2026M07")


def test_native_observation_ids_distinguish_government_airports():
    from src.transform.stage9_lineage import add_record_ids

    aicm, _ = native.aicm_pages(aicm_document(), original(), "2026M07")
    aifa, _ = native.aifa({"workbook": (workbook_bytes(aifa_document()), original("aifa"))}, "2026M07")
    ids = []
    for frame in [aicm, aifa]:
        subset = frame.loc[
            frame.period_id.eq("2026M07") & frame.metric_key.eq("cargo") & frame.market.eq("total")
        ]
        assert len(subset) == 1
        ids.append(
            add_record_ids(
                subset, table_name="government_native", natural_key_columns=native.GRAIN
            ).record_id.iloc[0]
        )
    assert ids[0] != ids[1]


def test_native_baseline_from_a_different_airport_is_not_a_revision(tmp_path):
    frame, _ = native.aicm_pages(aicm_document(), original(), "2026M07")
    before = frame.copy()
    before["airport_iata"] = "NLU"
    path = tmp_path / "baselines/other-airport.parquet"
    path.parent.mkdir()
    before.to_parquet(path, index=False)
    with pytest.raises(ValueError, match="same|airport"):
        candidate.revisions(
            tmp_path, "baselines/other-airport.parquet", "aicm_traffic", {"silver/government_native": frame}
        )


def test_aifa_native_baseline_includes_independent_ytd_values(tmp_path):
    frame, _ = native.aifa({"workbook": (workbook_bytes(aifa_document()), original("aifa"))}, "2026M07")
    before = frame.copy()
    key = before.index[
        before.period_type.eq("year_to_date") & before.metric_key.eq("cargo") & before.market.eq("total")
    ]
    assert len(key) == 1
    before.loc[key[0], "value"] += 0.1
    path = tmp_path / "baselines/native.parquet"
    path.parent.mkdir()
    before.to_parquet(path, index=False)
    result = candidate.revisions(
        tmp_path, "baselines/native.parquet", "aifa_traffic", {"silver/government_native": frame}
    )
    assert result["changed"] == 1

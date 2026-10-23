"""Strict native government-airport reports; never infer empty future months as zero."""

from __future__ import annotations

import calendar
import math
import re
from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import ZipFile

import pandas as pd
from openpyxl import load_workbook
from pypdf import PdfReader

from src.source_updates.calendar import local_time

MONTHS = "ENERO FEBRERO MARZO ABRIL MAYO JUNIO JULIO AGOSTO SEPTIEMBRE OCTUBRE NOVIEMBRE DICIEMBRE".split()
GRAIN = ["period_id", "period_type", "metric_key", "market"]
PRODUCTS = {"aicm_traffic", "aifa_traffic"}
VERSION = "government_airport_native_v1"


def month(period):
    if not re.fullmatch(r"20\d{2}M(?:0[1-9]|1[0-2])", period):
        raise ValueError("Government traffic needs an explicit report month, not a quarter")
    return int(period[:4]), int(period[-2:])


def number(value, *, decimals=False):
    if isinstance(value, bool) or value is None:
        raise ValueError("Missing or boolean native traffic value")
    text = str(value)
    pattern = r"(?:0|[1-9]\d*|[1-9]\d{0,2}(?:,\d{3})+)" + (r"\.\d" if decimals else "")
    if not re.fullmatch(pattern, text):
        raise ValueError("Native numeric grouping/precision changed")
    return Decimal(text.replace(",", ""))


def excel_number(value, *, mass=False):
    if type(value) not in {int, float}:
        raise ValueError("Missing, formula without cache or nonnumeric workbook value")
    try:
        result = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError("Invalid workbook number") from exc
    if not result.is_finite() or result < 0 or not mass and result != result.to_integral_value():
        raise ValueError("Invalid nonnegative native count/mass")
    return result


def reconcile(values, tolerance=Decimal(0)):
    if abs(values[0] + values[1] - values[2]) > tolerance:
        raise ValueError("Native domestic/international or arrival/departure sum differs")


def observation(original, period, kind, metric, market, value, unit, locator, scope):
    if not math.isfinite(float(value)) or abs(value) > 2**53:
        raise ValueError(
            "Native traffic magnitude cannot be represented exactly enough for candidate contracts"
        )
    return dict(
        period_id=period,
        period_type=kind,
        metric_key=metric,
        market=market,
        value=float(value),
        source_value=str(value),
        unit=unit,
        aviation_scope=scope,
        airport_iata="MEX" if original["source_system"] == "aicm" else "NLU",
        source_system=original["source_system"],
        source_file=original["source_file"],
        source_hash=original["sha256"],
        ingested_at=pd.Timestamp(original["downloaded_at"]).as_unit("ns"),
        parser_version=VERSION,
        source_locator=locator,
        agent_eligible=False,
    )


def aicm_pages(pages, original, period):
    year, end = month(period)
    if len(pages) != 14:
        raise ValueError("AICM reviewed report requires fourteen pages")
    cover = " ".join(pages[0].split())
    report = re.fullmatch(
        r"AICM en Cifras ([A-Za-z]+) (20\d{2}) Última actualización (\d{1,2}) de ([a-z]+) de (20\d{2})",
        cover,
    )
    if not report or report[1].upper() != MONTHS[end - 1] or int(report[2]) != year:
        raise ValueError("AICM cover does not identify the requested report month")
    if report[4].upper() not in MONTHS:
        raise ValueError("AICM update month unknown")
    updated = date(int(report[5]), MONTHS.index(report[4].upper()) + 1, int(report[3]))
    if (
        not date(year, end, calendar.monthrange(year, end)[1])
        <= updated
        <= local_time(original["downloaded_at"]).date()
    ):
        raise ValueError("AICM internal update date inconsistent with report/capture")
    rows, checks = [], 0
    for page, metric, unit, scope, title in [
        (2, "passengers", "passengers", "commercial", "Pasajeros Comerciales"),
        (7, "operations", "operations", "commercial_and_general", "Comerciales y de Aviación General"),
        (12, "cargo", "tonnes", "cargo_transported", "Carga Transportada"),
    ]:
        text = pages[page]
        native_unit = {"passengers": "Pasajeros", "operations": "Operaciones", "cargo": "Carga"}[metric]
        if title not in text or not re.search(
            rf"Enero-{MONTHS[end - 1].title()} {year}\s*[-–]\s*{native_unit}", text
        ):
            raise ValueError("AICM table scope/period header changed")
        if metric == "operations" and "Operaciones Totales" not in text:
            raise ValueError("AICM native total-operations scope changed")
        if metric == "cargo" and "en Toneladas" not in text:
            raise ValueError("AICM native cargo tonnes unit changed")
        if " ".join(["Nacional Internacional Total"] * 3) not in " ".join(text.split()):
            raise ValueError("AICM native market-column headings changed")
        header = (
            f"{year} Variación Porcentual{year - 1}"
            if metric == "cargo"
            else f"{year - 1} {year} Variación Porcentual"
        )
        if header not in text:
            raise ValueError("AICM comparative year header missing")
        decimals = metric == "cargo"
        # Rounded one-decimal components can differ by at most 0.15 tonnes.
        tolerance = Decimal("0.15") if decimals else Decimal(0)
        monthly = {year - 1: {}, year: {}}
        lines = text.splitlines()
        for m, label in enumerate(MONTHS, 1):
            candidates = [s.split() for s in lines if s.startswith(label + " ")]
            if len(candidates) != 1:
                raise ValueError("AICM native month missing or duplicated")
            tokens = candidates[0][1:]
            expected = 9 if m <= end else 3
            if len(tokens) != expected or any(not re.fullmatch(r"-?\d+\.\d", t) for t in tokens[6:]):
                raise ValueError("AICM monthly row incomplete or unexpected future observation")
            for current, offset in [(year - 1, 0)] + ([(year, 3)] if m <= end else []):
                values = [number(t, decimals=decimals) for t in tokens[offset : offset + 3]]
                reconcile(values, tolerance)
                checks += 1
                monthly[current][m] = values
                for c, (market, value) in enumerate(
                    zip(["domestic", "international", "total"], values, strict=True)
                ):
                    rows.append(
                        observation(
                            original,
                            f"{current}M{m:02d}",
                            "month",
                            metric,
                            market,
                            value,
                            unit,
                            f"page:{page + 1}/month:{label}/column:{offset + c + 1}",
                            scope,
                        )
                    )
        total_lines = [s for s in lines if s.startswith(f"Total {MONTHS[end - 1].title()} ")]
        annual = [s for s in lines if s.startswith("Total Anual ")]
        if len(total_lines) != 1 or len(annual) != 1:
            raise ValueError("AICM native accumulated/annual totals absent or duplicated")
        ytd_tokens, annual_tokens = total_lines[0].split()[2:], annual[0].split()[2:]
        if (
            len(ytd_tokens) != 9
            or len(annual_tokens) != 3
            or any(not re.fullmatch(r"-?\d+\.\d", t) for t in ytd_tokens[6:])
        ):
            raise ValueError("AICM accumulated/annual total format changed")
        for current, offset in [(year - 1, 0), (year, 3)]:
            values = [number(t, decimals=decimals) for t in ytd_tokens[offset : offset + 3]]
            reconcile(values, tolerance)
            checks += 1
            for c, (market, value) in enumerate(
                zip(["domestic", "international", "total"], values, strict=True)
            ):
                total = sum(monthly[current][m][c] for m in range(1, end + 1))
                if abs(total - value) > (end + 1) * Decimal("0.05") * decimals:
                    raise ValueError("AICM monthly observations do not reconcile to native YTD")
                checks += 1
                rows.append(
                    observation(
                        original,
                        f"{current}YTD{end:02d}",
                        "year_to_date",
                        metric,
                        market,
                        value,
                        unit,
                        f"page:{page + 1}/YTD/column:{offset + c + 1}",
                        scope,
                    )
                )
        for c, t in enumerate(annual_tokens):
            value = number(t, decimals=decimals)
            if abs(sum(monthly[year - 1][m][c] for m in range(1, 13)) - value) > Decimal("0.65") * decimals:
                raise ValueError("AICM previous-year annual reconciliation failed")
            checks += 1
    frame = pd.DataFrame(rows)
    return frame, dict(
        native_rows=len(frame),
        reconciliations=checks,
        report_period=period,
        months_validated=[f"{year}M{m:02d}" for m in range(1, end + 1)],
        source_update_date_as_reported=updated.isoformat(),
        first_publication_date=None,
        release_lag_days=None,
        scope="AICM_commercial_passengers_total_operations_and_cargo_not_airline_traffic",
        rounding_tolerance_tonnes="0.15_components; (months+1)*0.05_accumulation",
    )


def aicm(inputs, period):
    content, original = inputs["pdf"]
    if not content.startswith(b"%PDF-") or len(content) > 4_000_000:
        raise ValueError("AICM traffic PDF outside reviewed bound")
    return aicm_pages([p.extract_text() or "" for p in PdfReader(BytesIO(content)).pages], original, period)


def aifa(inputs, period):
    """Review only the retained 2026 workbook schema; contradictory dates fail closed."""
    content, original = inputs["workbook"]
    year, end = month(period)
    if year != 2026 or len(content) > 2_000_000 or not content.startswith(b"PK"):
        raise ValueError("AIFA workbook outside reviewed 2026 schema")
    with ZipFile(BytesIO(content)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 25_000_000:
            raise ValueError("AIFA expanded workbook exceeds reviewed bound")
    book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    if set(book.sheetnames) != {"Esquema Operativo", "Pasajeros", "Operaciones", "Carga", "Hoja1", "Tabs"}:
        raise ValueError("AIFA workbook sheets changed")
    rows, checks = [], 0
    accumulated = {}
    anchors = {
        "Pasajeros": {
            "C7": "AVIACIÓN COMERCIAL",
            "C9": "PASAJEROS COMERCIALES TOTALES.",
            "D10": "PASAJEROS",
            "D11": "PASAJEROS",
            "C68": "AVIACIÓN GENERAL",
            "C70": "PASAJEROS DE AVIACIÓN GENERAL.",
            "D72": "PASAJEROS",
        },
        "Operaciones": {
            "C7": "AVIACIÓN COMERCIAL",
            "C8": "OPERACIONES COMERCIALES",
            "D10": "OPERACIONES",
            "C66": "AVIACIÓN GENERAL",
            "C68": "OPERACIONES DE AVIACIÓN GENERAL",
            "D70": "OPERACIONES",
            "C96": "AVIACIÓN DE CARGA",
            "C97": "OPERACIONES DE CARGA",
            "D99": "OPERACIONES CON CARGA",
            "C128": "Nota 3: Las operaciones de Aviación de Carga se contabilizan sin aerolíneas mixtas.",
        },
        "Carga": {"C7": "AVIACIÓN DE CARGA", "C8": "CARGA TRANSPORTADA", "D10": "CARGA (KG.)"},
        "Esquema Operativo": {
            "E28": "Comercial",
            "I28": "Comercial",
            "E29": "General",
            "I29": "General",
            "E30": "Carga",
            "E9": "OPERACIONES",
            "I9": "PASAJEROS",
            "M9": "CARGA",
        },
    }
    for name, cells in anchors.items():
        if any(book[name][cell].value != text for cell, text in cells.items()):
            raise ValueError(f"AIFA {name}: aviation scope or native unit changed")
    # Dates are source metadata, never assumed first publication or repaired from filename.
    for name in ["Esquema Operativo", "Pasajeros", "Operaciones", "Carga"]:
        sheet = book[name]
        if not str(sheet["C2"].value).startswith(f"NUMERALIA AEROPORTUARIA {year}"):
            raise ValueError("AIFA workbook year/header changed")
        stamps = []
        for cell, prefix in [("C4", "Fecha de validación: "), ("C5", "Fecha de actualización: ")]:
            label = str(sheet[cell].value)
            match = re.fullmatch(re.escape(prefix) + r"(\d{1,2}) ([A-Z][a-z]{2})\. (20\d{2})", label)
            abbr = [m[:3].title() for m in MONTHS]
            if not match or match[2] not in abbr:
                raise ValueError("AIFA validation/update date missing or changed")
            stamp = date(int(match[3]), abbr.index(match[2]) + 1, int(match[1]))
            stamps.append(stamp)
            if (
                not date(year, end, calendar.monthrange(year, end)[1])
                <= stamp
                <= local_time(original["downloaded_at"]).date()
            ):
                raise ValueError(
                    f"AIFA {name}!{cell}: {label}; internal date precedes report month end or exceeds capture"
                )
        if stamps[0] > stamps[1]:
            raise ValueError("AIFA validation date exceeds its update date")
    for name, start, metric, unit, scope in [
        ("Pasajeros", 13, "passengers_commercial", "passengers", "commercial"),
        ("Pasajeros", 74, "passengers_general", "passengers", "general"),
        ("Operaciones", 12, "operations_commercial", "operations", "commercial"),
        ("Operaciones", 72, "operations_general", "operations", "general"),
        ("Operaciones", 101, "operations_cargo", "operations", "cargo_excludes_mixed_airlines"),
        ("Carga", 12, "cargo", "kg", "cargo_transported"),
    ]:
        sheet, values_by_month = book[name], {}
        mass = metric == "cargo"
        expected_header = (
            ["DE LLEGADA", "DE SALIDA", "SUBTOTAL"]
            if mass
            else [
                "NACIONAL" if metric == "operations_cargo" else "NACIONALES",
                "INTERNACIONAL" if metric == "operations_cargo" else "INTERNACIONALES",
                "SUBTOTAL",
            ]
        )
        if [sheet.cell(start - 1, c).value for c in [4, 5, 6]] != expected_header:
            raise ValueError("AIFA native market/unit header changed")
        if not any(
            f"{calendar.monthrange(year, end)[1]} {MONTHS[end - 1][:3].title()}. {year}"
            in str(sheet[f"C{r}"].value)
            for r in range(max(1, start - 5), start)
        ):
            raise ValueError("AIFA native block reporting window changed")
        if mass and sheet["D10"].value != "CARGA (KG.)":
            raise ValueError("AIFA native cargo unit changed")
        for m in range(1, 13):
            r = start + m - 1
            if str(sheet[f"C{r}"].value).rstrip("*") != MONTHS[m - 1]:
                raise ValueError("AIFA native monthly labels missing, duplicated or reordered")
            tokens = [sheet.cell(r, c).value for c in [4, 5, 6]]
            if m > end:
                if tokens != [None, None, 0]:
                    raise ValueError("AIFA future template row contains observations")
                continue
            values = [excel_number(t, mass=mass) for t in tokens]
            reconcile(values, Decimal("0.000001") if mass else Decimal(0))
            checks += 1
            values_by_month[m] = values
            for c, (market, value) in enumerate(
                zip(
                    ["arrival", "departure", "total"] if mass else ["domestic", "international", "total"],
                    values,
                    strict=True,
                ),
                4,
            ):
                rows.append(
                    observation(
                        original,
                        f"{year}M{m:02d}",
                        "month",
                        metric,
                        market,
                        value,
                        unit,
                        f"{name}!{chr(64 + c)}{r}",
                        scope,
                    )
                )
        total = start + 12
        if sheet[f"C{total}"].value != "TOTAL":
            raise ValueError("AIFA native annual-to-date total missing")
        for c in range(3):
            value = excel_number(sheet.cell(total, c + 4).value, mass=mass)
            if abs(sum(v[c] for v in values_by_month.values()) - value) > Decimal("0.000001") * mass:
                raise ValueError("AIFA monthly block/YTD mismatch")
            checks += 1
        accumulated[metric] = [sum(v[c] for v in values_by_month.values()) for c in range(3)]
    summary = book["Esquema Operativo"]
    if (
        summary["C28"].value
        != f"{year}\n(1/o. Ene. -{calendar.monthrange(year, end)[1]} {MONTHS[end - 1][:3].title()}.)"
        or summary["M11"].value != "Tons"
    ):
        raise ValueError("AIFA current-year summary window or tonnes unit changed")
    for cell, value, mass in [
        ("F28", accumulated["operations_commercial"][2], False),
        ("F29", accumulated["operations_general"][2], False),
        ("F30", accumulated["operations_cargo"][2], False),
        (
            "G28",
            sum(
                accumulated[k][2] for k in ["operations_commercial", "operations_general", "operations_cargo"]
            ),
            False,
        ),
        ("J28", accumulated["passengers_commercial"][2], False),
        ("J29", accumulated["passengers_general"][2], False),
        ("K28", sum(accumulated[k][2] for k in ["passengers_commercial", "passengers_general"]), False),
        ("M28", accumulated["cargo"][2] / 1000, True),
    ]:
        if abs(excel_number(summary[cell].value, mass=mass) - value) > Decimal("0.000001") * mass:
            raise ValueError("AIFA monthly blocks/whole-airport summary differ")
        checks += 1
    frame = pd.DataFrame(rows)
    return frame, dict(
        native_rows=len(frame),
        reconciliations=checks,
        report_period=period,
        months_validated=[f"{year}M{m:02d}" for m in range(1, end + 1)],
        scope="AIFA_commercial_general_cargo_separated_not_airline_traffic",
        first_publication_date=None,
        release_lag_days=None,
    )


def read(product, inputs, period):
    frame, details = {"aicm_traffic": aicm, "aifa_traffic": aifa}[product](inputs, period)
    if frame.empty or frame.duplicated(GRAIN).any() or frame.value.isna().any():
        raise ValueError("Native government report is empty, duplicated or missing")
    return frame, details

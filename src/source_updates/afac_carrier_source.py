"""Full monthly AFAC carrier dossier: bulletin counts, charter source and exact ancestry."""

from __future__ import annotations

import calendar
import csv
import hashlib
import io
import json
import re
from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from zipfile import ZipFile

import openpyxl
import pandas as pd
from pypdf import PdfReader

from src.source_updates.carrier_source_contract import clean, record
from src.source_updates.catalog import ROOT

PRODUCT = "afac_carriers"
REFERENCE = ROOT / "data/reference/carrier_crosswalk.csv"
HEADERS = ["Año", "Tipo", "Servicio", "Región", "Aerolinea", "Id_mes", "Mes", "Pasajeros", "Fecha"]
MONTHS = "enero febrero marzo abril mayo junio julio agosto septiembre octubre noviembre diciembre".split()
SHORT = "Ene Feb Mar Abr May Jun Jul Ago Sep Oct Nov Dic".split()
COUNT = r"(?:\d{1,3}(?:,\d{3})+|\d+)"
PRELIMINARY = (
    "Las cifras pueden variar debido a correcciones o ajustes por parte de la fuente, "
    "por lo que es importante tomarlas con carácter preliminar."
)
PERCENT = r"(?:-?\d+(?:\.\d+)?%|N\.C\.)"


def month(period):
    if not isinstance(period, str) or not re.fullmatch(r"2026M(?:0[1-9]|1[0-2])", period):
        raise ValueError("AFAC reviewed 2026 monthly bulletin/database format required")
    return int(period[-2:])


def end(year, m):
    return f"{year}-{m:02d}-{calendar.monthrange(year, m)[1]:02d}"


def mapping(name, stop):
    with REFERENCE.open(encoding="utf-8") as stream:
        rows = list(csv.DictReader(stream))
    matches = [
        (i + 2, r)
        for i, r in enumerate(rows)
        if r["source_system"] == "afac"
        and r["source_carrier_name"] == name
        and r["valid_from"] <= stop[:7]
        and (not r["valid_to"] or stop[:7] <= r["valid_to"])
    ]
    if len(matches) > 1:
        raise ValueError("Ambiguous effective AFAC carrier mapping; no first-match selection")
    if not matches:
        return None, None
    i, r = matches[0]
    evidence = json.dumps(
        dict(
            file="data/reference/carrier_crosswalk.csv",
            row=i,
            sha256=hashlib.sha256(REFERENCE.read_bytes()).hexdigest(),
        ),
        sort_keys=True,
    )
    return r["carrier_key"] or None, evidence


def pdf_values(line):
    parts = line.split()
    if len(parts) <= 12:
        raise ValueError("Incomplete AFAC source row")
    label, tokens = " ".join(parts[:-12]), parts[-12:]
    if not all(re.fullmatch(COUNT, tokens[i]) for i in [0, 1, 4, 5, 8, 9]) or not all(
        re.fullmatch(PERCENT, tokens[i]) for i in [2, 3, 6, 7, 10, 11]
    ):
        raise ValueError("AFAC row count/percent columns changed; never split numbers in carrier names")
    return label, [tokens[i] for i in [0, 1, 4, 5, 8, 9]]


def pdf_dates(column, m):
    year = 2024 if column == 1 else 2026 if column in [4, 6] else 2025
    kind = "year" if column <= 2 else "year_to_date" if column <= 4 else "month"
    stop = end(year, 12 if kind == "year" else m)
    start = f"{year}-{m:02d}-01" if kind == "month" else f"{year}-01-01"
    pid = (
        f"{year}FY"
        if kind == "year"
        else f"{year}YTD{m:02d}"
        if kind == "year_to_date"
        else f"{year}M{m:02d}"
    )
    return kind, pid, start, stop


def estimate(name, start, stop, notes):
    matches = re.findall(r"Datos estimados de (.+?) para ([a-z]+) (20\d{2})\.", notes)
    for fragment, month_name, year in matches:
        if month_name not in MONTHS:
            raise ValueError("Unreviewed AFAC estimate month")
        day = f"{year}-{MONTHS.index(month_name) + 1:02d}-01"
        if fragment.casefold() in name.casefold() and start <= day <= stop:
            return True
    return None  # No note is not evidence that a source cell is unestimated.


def bulletin(content, source, period):
    m = month(period)
    if len(content) > 3_000_000:
        raise ValueError("AFAC bulletin exceeds reviewed size")
    pages = PdfReader(io.BytesIO(content)).pages
    if len(pages) != 2:
        raise ValueError("AFAC scheduled bulletin must have exactly two pages")
    output, notes_by_market = [], {}
    for page, market in enumerate(["domestic", "international"], 1):
        text = pages[page - 1].extract_text()
        needed = f"y {MONTHS[m - 1]} de 2025 - 2026"
        normalized = " ".join(text.split()).replace("–", "-")
        if (
            needed not in normalized
            or "Notas:" not in text
            or "Fuente" not in text
            or ("vuelos nacionales" if market == "domestic" else "vuelos internacionales") not in normalized
            or "vuelos regulares" not in normalized
        ):
            raise ValueError("AFAC bulletin actual month, market, service or notes differ")
        notes = text.split("Notas:", 1)[1].split("Fuente", 1)[0].strip()
        if PRELIMINARY not in " ".join(notes.split()):
            raise ValueError("AFAC preliminary status needs its literal documentary declaration")
        notes_by_market[market] = notes
        data = text.split("Notas:", 1)[0].splitlines()
        if len(data) < 4 or "2024 2025" not in " ".join(data[:2]) or "2026" not in data[1]:
            raise ValueError("AFAC annual/YTD/month/current-comparative headers changed")
        group = "Total general" if market == "domestic" else None
        closed = False
        for line_number, line in enumerate(data[2:], 3):
            if not line.strip():
                continue
            if closed:
                raise ValueError("AFAC unexpected rows after the grand total")
            name, tokens = pdf_values(line)
            compact = re.sub(r"\s+", "", name).casefold()
            if compact in {"totalgeneral", "total"}:
                statement, row_group, closed = "pdf_total", "all_reported_carriers", True
            elif name.startswith("Total "):
                statement, group = "pdf_subtotal", name
                row_group = group
            else:
                if group is None:
                    raise ValueError("AFAC international operator precedes its literal subtotal group")
                statement, row_group = "pdf_operator", group
            for column, token in enumerate(tokens, 1):
                kind, p, start, stop = pdf_dates(column, m)
                carrier, mapped = mapping(name, stop) if statement == "pdf_operator" else (None, None)
                output.append(
                    record(
                        PRODUCT,
                        source,
                        edition=period,
                        entity=name,
                        carrier=carrier,
                        mapping=mapped,
                        statement=statement,
                        kind=kind,
                        period_id=p,
                        start=start,
                        end=stop,
                        metric="passengers_transported",
                        segment=market,
                        service="scheduled",
                        token=token,
                        unit="passengers_transported",
                        scale=1,
                        label=name,
                        locator=f"pdf/page[{page}]/line[{line_number}]/count[{column}]",
                        excerpt=line,
                        table=f"PDF:{page}",
                        row=line_number,
                        column=column,
                        notes=notes,
                        aviation="source_reported_scheduled_passengers",
                        group=row_group,
                        preliminary=True,
                        estimated=estimate(name, start, stop, notes),
                    )
                )
        if not closed:
            raise ValueError("AFAC grand total is missing")
    return output, notes_by_market


def database(content, source, period, notes):
    m = month(period)
    if len(content) > 16_000_000:
        raise ValueError("AFAC database ZIP exceeds reviewed bound")
    with ZipFile(io.BytesIO(content)) as archive:
        files = archive.infolist()
        if (
            len(files) != 1
            or files[0].filename != "DB_AFAC.xlsx"
            or files[0].flag_bits & 1
            or (files[0].file_size > 10_000_000 or (files[0].external_attr >> 16) & 0o170000 == 0o120000)
        ):
            raise ValueError("AFAC ZIP needs its unique bounded ordinary DB_AFAC.xlsx member")
        body = archive.read(files[0])
    with ZipFile(io.BytesIO(body)) as archive:
        files = archive.infolist()
        if (
            len(files) > 1000
            or len({f.filename for f in files}) != len(files)
            or sum(f.file_size for f in files) > 100_000_000
            or any("externalLink" in f.filename or f.flag_bits & 1 for f in files)
        ):
            raise ValueError("AFAC workbook duplicate/oversized/encrypted/external-linked members")
        if any(
            b"<!ENTITY" in archive.read(f).upper() or b"<!DOCTYPE" in archive.read(f).upper()
            for f in files
            if f.filename.endswith(".xml")
        ):
            raise ValueError("AFAC workbook custom XML declarations are forbidden")
    book = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=False)
    output, covered = [], defaultdict(set)
    try:
        if book.sheetnames != ["AFAC"]:
            raise ValueError("AFAC workbook sheet inventory changed")
        rows = book["AFAC"].iter_rows(values_only=True)
        if list(next(rows)) != HEADERS:
            raise ValueError("AFAC database columns changed")
        for i, row in enumerate(rows, 2):
            if i > 100_000:
                raise ValueError("AFAC database row count exceeds bound")
            if all(v is None for v in row):
                continue  # Formatting-only blank rows are not source observations.
            year, market, service, region, name, mo, short, value, date = row
            if type(year) is not int or not 1900 <= year <= 2026:
                raise ValueError("Invalid AFAC database year")
            if year not in {2024, 2025, 2026}:
                continue
            market, service, region, name, short = [
                v.strip() if isinstance(v, str) else v for v in [market, service, region, name, short]
            ]
            if (
                type(mo) is not int
                or not 1 <= mo <= 12
                or short != SHORT[mo - 1]
                or market not in {"Nacional", "Internacional"}
                or service not in {"Regular", "Fletamento"}
                or not all(isinstance(v, str) and v.strip() == v and v for v in [region, name])
                or not isinstance(date, datetime)
                or date != datetime(year, mo, 1)
                or type(value) not in {int, float}
                or not Decimal(str(value)).is_finite()
                or value < 0
                or int(value) != value
                or value > 2**53 - 1
            ):
                raise ValueError(
                    f"AFAC database identity, date, service or passenger count changed at row {i}"
                )
            covered[year, service].add(mo)
            start, stop = f"{year}-{mo:02d}-01", end(year, mo)
            mapped, evidence = mapping(name, stop)
            segment = "domestic" if market == "Nacional" else "international"
            output.append(
                record(
                    PRODUCT,
                    source,
                    edition=period,
                    entity=name,
                    carrier=mapped,
                    mapping=evidence,
                    statement="database_row",
                    kind="month",
                    period_id=f"{year}M{mo:02d}",
                    start=start,
                    end=stop,
                    metric="passengers_transported",
                    segment=segment,
                    service="scheduled" if service == "Regular" else "charter",
                    token=str(int(value)),
                    unit="passengers_transported",
                    scale=1,
                    label=name,
                    locator=f"zip:DB_AFAC.xlsx/AFAC!R{i}C8",
                    excerpt=json.dumps(
                        [v.isoformat() if isinstance(v, datetime) else v for v in row], ensure_ascii=False
                    ),
                    table="DB_AFAC.xlsx/AFAC",
                    row=i,
                    column=8,
                    notes=(
                        "Database row; surrounding label whitespace trimmed, original tokens in excerpt. "
                        "No independent release date or blanket estimation/preliminary flag."
                    ),
                    aviation="source_reported_carrier_passengers",
                    group=region,
                    estimated=None,
                )
            )
        for year in [2024, 2025, 2026]:
            for service in ["Regular", "Fletamento"]:
                if covered[year, service] != set(range(1, (m if year == 2026 else 12) + 1)):
                    raise ValueError("AFAC database month/service coverage does not match bulletin edition")
    finally:
        book.close()
    return output, hashlib.sha256(body).hexdigest()


def validate(frame):
    if frame.empty or frame.edition.nunique() != 1:
        raise ValueError("Complete single-edition AFAC dossier required")
    period = frame.edition.iloc[0]
    month(period)
    if set(frame.statement) != {"pdf_operator", "pdf_subtotal", "pdf_total", "database_row"}:
        raise ValueError("AFAC source table coverage differs")
    if frame.duplicated(["source_hash", "source_locator"]).any():
        raise ValueError("Duplicate AFAC physical source cell")
    operators = frame.loc[frame.statement.eq("pdf_operator")]
    if operators.duplicated(["period_id", "entity_as_reported", "segment"]).any():
        raise ValueError("AFAC PDF repeats an operator/month or aggregate observation")
    charter = frame.loc[
        frame.statement.eq("database_row") & frame.service_type.eq("charter") & frame.period_id.eq(period)
    ]
    if charter.duplicated(["entity_as_reported", "segment", "service_type"]).any():
        raise ValueError("Ambiguous current AFAC charter duplicate; no implicit summation/deduplication")
    for raw in frame.to_dict("records"):
        validate_row(clean(raw))
    return frame


def validate_row(r):
    """Bind each row, including baseline/Gold subsets, to reviewed source semantics."""
    from types import SimpleNamespace

    r = SimpleNamespace(**r)
    period = r.edition
    m = month(period)
    if (r.unit, r.scale, r.metric_key) != ("passengers_transported", 1.0, "passengers_transported") or (
        pd.isna(r.value) or r.value < 0 or r.value % 1 or r.value > 2**53 - 1
    ):
        raise ValueError("AFAC verified passenger counts and units required")
    if r.segment not in {"domestic", "international"} or r.source_system != "afac" or not r.source_group:
        raise ValueError("AFAC market, source alias or literal group differs")
    if r.statement.startswith("pdf_"):
        page = 1 if r.segment == "domestic" else 2
        if r.statement not in {"pdf_operator", "pdf_subtotal", "pdf_total"} or not 1 <= r.source_column <= 6:
            raise ValueError("AFAC reviewed PDF row/column required")
        if (r.period_type, r.period_id, r.period_start_date, r.period_end_date) != pdf_dates(
            r.source_column, m
        ):
            raise ValueError("AFAC annual/YTD/month semantics changed")
        if (
            r.service_type != "scheduled"
            or r.is_preliminary is not True
            or (PRELIMINARY not in " ".join(r.source_notes.split()))
        ):
            raise ValueError("AFAC PDF service/preliminary documentary semantics changed")
        if (
            r.source_table != f"PDF:{page}"
            or r.source_locator != (f"pdf/page[{page}]/line[{r.source_row_number}]/count[{r.source_column}]")
            or r.source_row_number < 3
            or r.aviation_scope != "source_reported_scheduled_passengers"
        ):
            raise ValueError("AFAC PDF physical locator or scope differs")
        if r.source_estimated != estimate(
            r.entity_as_reported, r.period_start_date, r.period_end_date, r.source_notes
        ):
            raise ValueError("AFAC estimate flag is not backed by the cell's own source note")
    elif r.statement == "database_row":
        if (
            not re.fullmatch(r"202[456]M(?:0[1-9]|1[0-2])", r.period_id)
            or r.period_type != "month"
            or r.service_type not in {"scheduled", "charter"}
        ):
            raise ValueError("AFAC database monthly service/year differs")
        year, mo = int(r.period_id[:4]), int(r.period_id[-2:])
        if (
            year == 2026
            and mo > m
            or r.period_start_date != f"{year}-{mo:02d}-01"
            or r.period_end_date != end(year, mo)
        ):
            raise ValueError("AFAC database month dates/edition differ")
        if r.is_preliminary is not None or r.source_estimated is not None:
            raise ValueError("AFAC database cannot borrow PDF preliminary/estimated declarations")
        if (
            r.source_table != "DB_AFAC.xlsx/AFAC"
            or r.source_column != 8
            or r.source_row_number < 2
            or (r.source_locator != f"zip:DB_AFAC.xlsx/AFAC!R{r.source_row_number}C8")
            or r.aviation_scope != "source_reported_carrier_passengers"
        ):
            raise ValueError("AFAC database physical locator/scope differs")
    else:
        raise ValueError("AFAC source statement differs")
    if r.is_unaudited is not None or r.source_label != r.entity_as_reported:
        raise ValueError("AFAC has no financial audit declaration or renamed source entity")
    expected = (
        mapping(r.entity_as_reported, r.period_end_date)
        if r.statement in {"database_row", "pdf_operator"}
        else (None, None)
    )
    if (r.carrier_key, r.mapping_evidence) != expected:
        raise ValueError("AFAC entity mapping is not the exact versioned effective source name")


def reconcile(frame):
    validate(frame)
    db = frame.loc[frame.statement.eq("database_row")]
    pdf = frame.loc[frame.statement.ne("database_row")]
    operators = pdf.loc[pdf.statement.eq("pdf_operator")]
    checks, explicit_zero_without_db = 0, []
    for r in pdf.loc[pdf.statement.isin(["pdf_total", "pdf_subtotal"])].itertuples(index=False):
        selected = operators.loc[operators.segment.eq(r.segment) & operators.period_id.eq(r.period_id)]
        if r.statement == "pdf_subtotal":
            selected = selected.loc[selected.source_group.eq(r.source_group)]
        if selected.empty or selected.value.sum() != r.value:
            raise ValueError("AFAC exact PDF regional/grand total failed")
        checks += 1
    for r in operators.itertuples(index=False):
        selected = db.loc[
            db.entity_as_reported.eq(r.entity_as_reported)
            & db.segment.eq(r.segment)
            & db.service_type.eq("scheduled")
            & db.period_start_date.ge(r.period_start_date)
            & db.period_end_date.le(r.period_end_date)
        ]
        if selected.empty and r.value == 0:
            explicit_zero_without_db.append(r.source_locator)
        elif selected.empty or selected.value.sum() != r.value:
            raise ValueError("AFAC database versus bulletin annual/YTD/month cell differs")
        else:
            checks += 1
    regular = db.loc[db.service_type.eq("scheduled")]
    names = set(operators[["entity_as_reported", "segment"]].itertuples(index=False, name=None))
    if any(
        (r.entity_as_reported, r.segment) not in names and r.value != 0
        for r in regular.itertuples(index=False)
    ):
        raise ValueError("AFAC database has positive regular operators missing from bulletin")
    return dict(
        reconciliations_passed=checks,
        exact_count_reconciliation=True,
        explicit_PDF_zero_without_database_observation=explicit_zero_without_db,
        source_estimated_observations=int(frame.source_estimated.eq(True).sum()),
        unmapped_native_rows=int(frame.carrier_key.isna().sum()),
        no_operator_route_allocation=True,
        no_database_absence_imputed_as_zero=True,
    )


def read(inputs, period):
    m = month(period)
    if set(inputs) != {"pdf", "database"}:
        raise ValueError("AFAC monthly dossier requires bulletin PDF and database ZIP")
    urls = dict(
        pdf=f"https://datatur.sectur.gob.mx/Documentoscompartidos/afac/AFAC_2026_{m:02d}.pdf",
        database="https://datatur.sectur.gob.mx/Documentoscompartidos/afac/DB_AFAC.zip",
    )
    for role, (_, source) in inputs.items():
        if (
            source["source_system"] != "afac"
            or source["source_url"] != urls[role]
            or pd.Timestamp(source["downloaded_at"]).date().isoformat() <= end(2026, m)
        ):
            raise ValueError("AFAC original URL/alias/capture does not identify a closed monthly dossier")
    a, notes = bulletin(*inputs["pdf"], period)
    b, member_hash = database(*inputs["database"], period, notes)
    frame = pd.DataFrame(a + b)
    details = reconcile(frame)
    return frame, {
        **details,
        "scope": "retained_AFAC_monthly_carrier_passengers_regular_and_charter_with_source_caveats",
        "database_member_sha256": member_hash,
        "source_notes_by_market": notes,
        "monthly_coverage_verified": True,
        "complete_month_certified": False,
        "source_declares_incomplete_international_reporting": "no incluyen la totalidad"
        in notes["international"],
    }


def project(native):
    reconcile(native)
    period = native.edition.iloc[0]
    regular = native.loc[native.statement.eq("pdf_operator") & native.period_id.eq(period)].copy()
    result = []
    for _, r in regular.iterrows():
        match = native.loc[
            native.statement.eq("database_row")
            & native.service_type.eq("scheduled")
            & native.period_id.eq(period)
            & native.segment.eq(r.segment)
            & native.entity_as_reported.eq(r.entity_as_reported),
            "record_id",
        ]
        r["_parent_ids"] = [r.record_id, *sorted(match)]
        result.append(r)
    for _, r in native.loc[
        native.statement.eq("database_row") & native.service_type.eq("charter") & native.period_id.eq(period)
    ].iterrows():
        r["_parent_ids"] = [r.record_id]
        result.append(r)
    return pd.DataFrame(result)

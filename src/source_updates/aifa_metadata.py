"""Bounded metadata for the reviewed September 2026 AIFA workbook edition."""

from io import BytesIO

from openpyxl import load_workbook

NOTES = (
    "*Por ratificación del área operativa se actualizaron las cifras de la Aviación General "
    "de Junio de 2026 ( 1 operación y 0 pasajeros).",
    "**Por ratificación del área operativa se actualizaron las cifras de la Aviación General "
    "de Agosto de 2026 ( 3 operaciones y 4 pasajeros).",
)


def september_metadata(content, book, period, expected_window):
    """Return metadata only for this reviewed layout; never invent revision deltas."""
    if period != "2026M09":
        raise ValueError("AIFA shifted layout is reviewed only for September 2026")
    formulas = load_workbook(BytesIO(content), read_only=True, data_only=False)
    try:
        references = {
            ("Esquema Operativo", "C4"): "=Pasajeros!C4",
            ("Esquema Operativo", "C5"): "=Pasajeros!C5",
            ("Operaciones", "C4"): "=Pasajeros!C4",
            ("Operaciones", "C5"): "=Pasajeros!C5",
            ("Carga", "C4"): "=Pasajeros!C4",
            ("Carga", "C5"): "=Pasajeros!C10",
            ("Carga", "C9"): "=C5",
        }
        for (sheet, cell), reference in references.items():
            cached = book[sheet][cell].value
            if formulas[sheet][cell].value != reference:
                raise ValueError("AIFA September metadata formula changed")
            target_sheet, target_cell = (
                reference[1:].split("!") if "!" in reference else (sheet, reference[1:])
            )
            if cached != book[target_sheet][target_cell].value:
                raise ValueError("AIFA September metadata cache differs from its referenced cell")
        if book["Carga"]["C5"].value != expected_window:
            raise ValueError("AIFA cargo metadata must preserve its reporting window, not an inferred date")
        notes = []
        for text, period_id, pax_cell, ops_cell in zip(
            NOTES, ("2026M06", "2026M08"), ("C100", "C101"), ("C95", "C96"), strict=True
        ):
            if book["Pasajeros"][pax_cell].value != text or book["Operaciones"][ops_cell].value != text:
                raise ValueError("AIFA September revision note missing or changed")
            notes.append(
                {
                    "text": text,
                    "period_id": period_id,
                    "locators": [f"Pasajeros!{pax_cell}", f"Operaciones!{ops_cell}"],
                    "reported_adjustment_delta": None,
                    "interpretation": "source_reports_revision_without_signed_delta_or_comparison_baseline",
                }
            )
        for sheet, start in (("Pasajeros", 74), ("Operaciones", 72)):
            for month in range(1, 13):
                marker = "**" if month == 8 else "*" if month == 6 else ""
                label = str(book[sheet].cell(start + month - 1, 3).value)
                if label != label.rstrip("*") + marker:
                    raise ValueError("AIFA revision markers differ from their reviewed monthly notes")
        return {"reported_revision_notes": notes, "layout": "aifa_2026_september_cargo_offset_2"}
    finally:
        formulas.close()

"""Attribute two AIFA domestic markets to Aeromexico with explicit inference.

AFAC reports airport-city pairs and movements, not carriers. A dated airline
roster names Aeromexico as the sole listed carrier on the Colima and Durango
AIFA routes. Keep that attribution distinct from observed carrier-route data.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
import unicodedata

from bs4 import BeautifulSoup
import duckdb
import openpyxl
import pandas as pd

from src.config import PATHS
from src.ingest.stage4_common import write_parquet_atomic
from src.transform.stage9_lineage import (
    LineageSpec, build_bridge_record_lineage, build_dim_source,
    build_dim_source_artifact, make_record_id,
)


AFAC = PATHS.bronze / "afac_research" / "afac_research_city_pairs_2026M07_20260908T182059Z.xlsx"
ROSTER = PATHS.bronze / "domestic_routes_research" / "expansion_aifa_airline_destinations_2026M06_20260913T200713Z.html"
TRANSFER = PATHS.bronze / "domestic_routes_research" / "colima_subsectur_mirror_colima_aicm_transfer_2026M05_20260913T200338Z.html"
SILVER = PATHS.silver / "domestic_routes" / "afac_exclusive_market_months.parquet"
GOLD = PATHS.gold / "fact_domestic_exclusive_market_inferences.parquet"
QUALITY = PATHS.quality / "afac_exclusive_domestic.json"
TABLE = "fact_domestic_exclusive_market_inferences"
BRIDGE = "bridge_domestic_exclusive_market_lineage"
VERSION = "afac_exclusive_domestic_v1"
CITY_TO_AIRPORT = {"SANTA LUCIA": "NLU", "MEXICO": "MEX", "COLIMA": "CLQ", "DURANGO": "DGO"}
REQUIRED_PAIRS = {
    ("SANTA LUCIA", "COLIMA"), ("COLIMA", "SANTA LUCIA"),
    ("MEXICO", "COLIMA"), ("COLIMA", "MEXICO"),
    ("SANTA LUCIA", "DURANGO"), ("DURANGO", "SANTA LUCIA"),
}


def _verified(path: Path) -> tuple[str, dict]:
    metadata = json.loads(Path(str(path) + ".meta.json").read_text(encoding="utf-8"))
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if digest != metadata["sha256"]:
        raise ValueError(f"Bronze hash mismatch: {path}")
    return digest, metadata


def _roster() -> dict[str, list[str]]:
    soup = BeautifulSoup(ROSTER.read_bytes().decode("utf-8"), "html.parser")
    sections: dict[str, list[str]] = {}
    for heading in soup.find_all("h3")[:5]:
        title = _ascii(heading.get_text(" ", strip=True)).lower()
        parts = []
        node = heading.next_sibling
        while node is not None and getattr(node, "name", None) not in {"h2", "h3"}:
            value = node.get_text(" ", strip=True) if hasattr(node, "get_text") else str(node).strip()
            if value:
                parts.append(value)
            node = node.next_sibling
        sections[title] = parts
    if set(sections) != {"aeromexico", "aerus", "mexicana", "viva", "volaris"}:
        raise ValueError("Unexpected airline sections in dated AIFA roster")
    for city in ("Colima", "Durango"):
        if city not in sections["aeromexico"]:
            raise ValueError(f"Aeromexico {city} route missing from dated roster")
        for carrier, destinations in sections.items():
            if carrier != "aeromexico" and city in destinations:
                raise ValueError(f"{city} is not exclusive in dated roster")
    return sections


def _ascii(value: str) -> str:
    return unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")


def _city(value: object) -> str:
    label = _ascii(str(value or "")).upper()
    if label.startswith("SANTA LUC"):
        return "SANTA LUCIA"  # AFAC XLSX includes a corrupted accented vowel.
    return label


def run() -> dict:
    afac_hash, afac_meta = _verified(AFAC)
    roster_hash, roster_meta = _verified(ROSTER)
    transfer_hash, transfer_meta = _verified(TRANSFER)
    _roster()
    transfer = BeautifulSoup(TRANSFER.read_bytes(), "html.parser").get_text(" ", strip=True).lower()
    if not all(token in transfer for token in ("25 de junio", "aifa", "aicm", "colima")):
        raise ValueError("Colima transfer date or airport missing")

    sheet = openpyxl.load_workbook(AFAC, read_only=True, data_only=True)["REG NAC"]
    if [sheet.cell(6, col).value for col in (6, 7, 8)] != ["Abr/Apr", "May/May", "Jun/Jun"]:
        raise ValueError("AFAC month columns changed")
    rows = []
    found = set()
    for row_number, cells in enumerate(sheet.iter_rows(min_row=7, values_only=True), start=7):
        pair = (_city(cells[0]), _city(cells[1]))
        if pair not in REQUIRED_PAIRS:
            continue
        if pair in found:
            raise ValueError(f"Duplicate AFAC market pair: {pair}")
        found.add(pair)
        for month, column in ((4, 6), (5, 7), (6, 8)):
            movements = cells[column - 1]
            if not isinstance(movements, (int, float)) or movements < 0:
                raise ValueError(f"Invalid AFAC movements at {column}{row_number}")
            if movements == 0:
                continue
            # The Colima press announcement provides the June 25 airport split.
            if pair in {("MEXICO", "COLIMA"), ("COLIMA", "MEXICO")} and (month != 6 or movements != 6):
                raise ValueError("MEX-Colima differs from six post-transfer June days")
            if pair in {("SANTA LUCIA", "COLIMA"), ("COLIMA", "SANTA LUCIA")} and month == 6 and movements != 24:
                raise ValueError("NLU-Colima differs from 24 pre-transfer June days")
            rows.append({
                "period_id": f"2026M{month:02d}",
                "origin_iata": CITY_TO_AIRPORT[pair[0]],
                "dest_iata": CITY_TO_AIRPORT[pair[1]],
                "market_movements": int(movements),
                "attribution_status": "inferred_exclusive_carrier_market",
                "carrier_group": "Aeromexico_and_Connect_unresolved",
                "source_sheet": "REG NAC", "source_cell": f"{openpyxl.utils.get_column_letter(column)}{row_number}",
                "afac_source_hash": afac_hash, "roster_source_hash": roster_hash,
                "transfer_source_hash": transfer_hash if "CLQ" in (CITY_TO_AIRPORT[pair[0]], CITY_TO_AIRPORT[pair[1]]) else None,
                "source_url": afac_meta["source_url"], "parser_version": VERSION,
            })
    if found != REQUIRED_PAIRS:
        raise ValueError(f"Missing AFAC market pairs: {REQUIRED_PAIRS - found}")
    silver = pd.DataFrame(rows).sort_values(["period_id", "origin_iata", "dest_iata"]).reset_index(drop=True)
    if silver.duplicated(["period_id", "origin_iata", "dest_iata"]).any():
        raise ValueError("Duplicate AFAC exclusive-market key")
    expected = {"CLQ": 182, "DGO": 182}
    for airport, total in expected.items():
        actual = int(silver[(silver.origin_iata.eq(airport)) | (silver.dest_iata.eq(airport))].market_movements.sum())
        if actual != total:
            raise ValueError(f"Unexpected {airport} quarterly movements: {actual}")
    SILVER.parent.mkdir(parents=True, exist_ok=True)
    write_parquet_atomic(silver, SILVER)
    gold = silver.copy()
    gold["source_system"] = "afac_od_exclusive_market_inference"
    gold["agent_eligible"] = False
    gold["record_id"] = gold.apply(lambda row: make_record_id(TABLE, {
        "afac_hash": afac_hash, "roster_hash": roster_hash,
        "period_id": row.period_id, "origin_iata": row.origin_iata,
        "dest_iata": row.dest_iata, "parser_version": VERSION,
    }), axis=1)
    artifacts = build_dim_source_artifact()
    artifact_ids = {}
    for name, digest in (("afac", afac_hash), ("roster", roster_hash), ("transfer", transfer_hash)):
        matched = artifacts[artifacts.artifact_sha256.eq(digest)]
        if len(matched) != 1:
            raise ValueError(f"{name} original does not resolve to one Bronze artifact")
        artifact_ids[name] = str(matched.iloc[0].artifact_id)
    gold["artifact_id"] = artifact_ids["afac"]
    bridge = build_bridge_record_lineage([
        LineageSpec(record_id=row.record_id, table_name=TABLE, lineage_type="derived",
                    artifact_ids=tuple(artifact_ids[name] for name in (("afac", "roster", "transfer") if "CLQ" in (row.origin_iata, row.dest_iata) else ("afac", "roster"))),
                    lineage_note=f"AFAC {row.source_sheet}!{row.source_cell}; carrier attributed from June 8 AIFA roster, not present in AFAC row")
        for row in gold.itertuples(index=False)
    ], artifacts)
    outputs = {TABLE: gold, BRIDGE: bridge, "dim_source": build_dim_source(),
               "dim_source_artifact": artifacts}
    with duckdb.connect(str(PATHS.warehouse)) as connection:
        connection.execute("BEGIN TRANSACTION")
        for name, frame in outputs.items():
            path = PATHS.gold / f"{name}.parquet"
            write_parquet_atomic(frame, path)
            connection.execute(f"CREATE OR REPLACE TABLE {name} AS SELECT * FROM read_parquet(?)", [str(path)])
        connection.execute("COMMIT")
    report = {
        "source_afac": str(AFAC.relative_to(PATHS.root)).replace("\\", "/"),
        "source_roster": str(ROSTER.relative_to(PATHS.root)).replace("\\", "/"),
        "source_transfer": str(TRANSFER.relative_to(PATHS.root)).replace("\\", "/"),
        "source_hashes": {"afac": afac_hash, "roster": roster_hash, "transfer": transfer_hash},
        "silver_rows": len(silver), "gold_rows": len(gold), "lineage_links": len(bridge),
        "quarter_movements_by_counterpart": expected,
        "quarter_market_movements": int(gold.market_movements.sum()),
        "attribution_status": "inferred_exclusive_carrier_market",
        "agent_eligible": False,
    }
    QUALITY.parent.mkdir(parents=True, exist_ok=True)
    QUALITY.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return report


if __name__ == "__main__":
    print(json.dumps(run(), ensure_ascii=False, indent=2))

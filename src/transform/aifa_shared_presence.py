"""Document Aeromexico's dated AIFA route presence without assigning market totals.

The June 8 airline roster names the carrier on eight AIFA routes also served by
other airlines. AFAC supplies quarterly city-pair activity, but no carrier split.
"""

from __future__ import annotations

import json

import duckdb
import openpyxl
import pandas as pd

from src.config import PATHS
from src.ingest.stage4_common import write_parquet_atomic
from src.transform.afac_exclusive_domestic import AFAC, ROSTER, _ascii, _city, _roster, _verified
from src.transform.stage9_lineage import (
    LineageSpec, build_bridge_record_lineage, build_dim_source,
    build_dim_source_artifact, make_record_id,
)


DESTINATIONS = {
    "CANCUN": "CUN", "GUADALAJARA": "GDL", "MERIDA": "MID",
    "MONTERREY": "MTY", "OAXACA": "OAX", "PUERTO VALLARTA": "PVR",
    "TULUM": "TQO", "VERACRUZ": "VER",
}
SILVER = PATHS.silver / "domestic_routes" / "aifa_shared_market_activity.parquet"
GOLD = PATHS.gold / "fact_aifa_shared_route_presence.parquet"
QUALITY = PATHS.quality / "aifa_shared_presence.json"
TABLE = "fact_aifa_shared_route_presence"
BRIDGE = "bridge_aifa_shared_route_presence_lineage"
VERSION = "aifa_shared_presence_v1"


def run() -> dict:
    afac_hash, afac_meta = _verified(AFAC)
    roster_hash, roster_meta = _verified(ROSTER)
    roster = _roster()
    aeromexico_text = _ascii(" ".join(roster["aeromexico"])).upper()
    for city in DESTINATIONS:
        if city not in aeromexico_text:
            raise ValueError(f"AIFA roster no longer identifies Aeromexico to {city}")
        if not any(city in _ascii(" ".join(values)).upper()
                   for carrier, values in roster.items() if carrier != "aeromexico"):
            raise ValueError(f"Shared-carrier status not supported for {city}")

    sheet = openpyxl.load_workbook(AFAC, read_only=True, data_only=True)["REG NAC"]
    if [sheet.cell(6, column).value for column in (6, 7, 8)] != ["Abr/Apr", "May/May", "Jun/Jun"]:
        raise ValueError("AFAC month columns changed")
    needed = {("SANTA LUCIA", city) for city in DESTINATIONS} | {
        (city, "SANTA LUCIA") for city in DESTINATIONS
    }
    found: set[tuple[str, str]] = set()
    rows: list[dict] = []
    for row_number, cells in enumerate(sheet.iter_rows(min_row=7, values_only=True), start=7):
        pair = (_city(cells[0]), _city(cells[1]))
        if pair not in needed:
            continue
        if pair in found:
            raise ValueError(f"Duplicate AFAC shared market pair: {pair}")
        found.add(pair)
        counterpart = pair[1] if pair[0] == "SANTA LUCIA" else pair[0]
        for month, column in ((4, 6), (5, 7), (6, 8)):
            value = cells[column - 1]
            if not isinstance(value, (int, float)) or value <= 0 or int(value) != value:
                raise ValueError(f"Invalid AFAC shared-market activity at {column}{row_number}")
            rows.append({
                "period_id": f"2026M{month:02d}",
                "origin_iata": "NLU" if pair[0] == "SANTA LUCIA" else DESTINATIONS[counterpart],
                "dest_iata": "NLU" if pair[1] == "SANTA LUCIA" else DESTINATIONS[counterpart],
                "market_movements_all_carriers": int(value),
                "source_sheet": "REG NAC",
                "source_cell": f"{openpyxl.utils.get_column_letter(column)}{row_number}",
                "afac_source_hash": afac_hash,
                "parser_version": VERSION,
            })
    if found != needed:
        raise ValueError(f"Missing AFAC shared markets: {needed - found}")
    silver = pd.DataFrame(rows).sort_values(["period_id", "origin_iata", "dest_iata"]).reset_index(drop=True)
    if silver.duplicated(["period_id", "origin_iata", "dest_iata"]).any() or len(silver) != 48:
        raise ValueError("Unexpected AIFA shared-market grain")

    gold_rows = []
    for city, airport in DESTINATIONS.items():
        part = silver[(silver.origin_iata.eq(airport)) | (silver.dest_iata.eq(airport))]
        if len(part) != 6:
            raise ValueError(f"Incomplete AFAC market context for {airport}")
        market_key = "<>".join(sorted(("NLU", airport)))
        gold_rows.append({
            "period_id": "2026Q2",
            "market_key": market_key,
            "origin_iata": "NLU",
            "dest_iata": airport,
            "carrier_group": "Aeromexico_or_Connect_unresolved",
            "carrier_presence_as_of": "2026-06-08",
            "carrier_movements": None,
            "market_movements_all_carriers": int(part.market_movements_all_carriers.sum()),
            "attribution_status": "carrier_route_present_volume_unresolved",
            "source_cells": json.dumps(sorted(part.source_cell.tolist())),
            "afac_source_hash": afac_hash,
            "roster_source_hash": roster_hash,
            "afac_source_url": afac_meta["source_url"],
            "roster_source_url": roster_meta["source_url"],
            "agent_eligible": False,
            "parser_version": VERSION,
            "record_id": make_record_id(TABLE, {
                "period_id": "2026Q2", "market_key": market_key,
                "afac_hash": afac_hash, "roster_hash": roster_hash,
                "parser_version": VERSION,
            }),
        })
    gold = pd.DataFrame(gold_rows).sort_values("market_key").reset_index(drop=True)
    if gold.duplicated(["period_id", "market_key"]).any():
        raise ValueError("Duplicate AIFA shared-presence key")

    artifacts = build_dim_source_artifact()
    artifact_ids = {}
    for name, digest in (("afac", afac_hash), ("roster", roster_hash)):
        matched = artifacts[artifacts.artifact_sha256.eq(digest)]
        if len(matched) != 1:
            raise ValueError(f"{name} original does not resolve to one Bronze artifact")
        artifact_ids[name] = str(matched.iloc[0].artifact_id)
    gold["artifact_id"] = artifact_ids["roster"]
    bridge = build_bridge_record_lineage([
        LineageSpec(
            record_id=row.record_id, table_name=TABLE, lineage_type="derived",
            artifact_ids=(artifact_ids["afac"], artifact_ids["roster"]),
            lineage_note=f"June 8 carrier route roster; AFAC {row.source_cells} is all-carrier market activity, not Aeromexico volume",
        ) for row in gold.itertuples(index=False)
    ], artifacts)

    SILVER.parent.mkdir(parents=True, exist_ok=True)
    write_parquet_atomic(silver, SILVER)
    outputs = {TABLE: gold, BRIDGE: bridge, "dim_source": build_dim_source(),
               "dim_source_artifact": artifacts}
    with duckdb.connect(str(PATHS.warehouse)) as connection:
        connection.execute("BEGIN TRANSACTION")
        for name, frame in outputs.items():
            path = PATHS.gold / f"{name}.parquet"
            write_parquet_atomic(frame, path)
            connection.execute(f"CREATE OR REPLACE TABLE {name} AS SELECT * FROM read_parquet(?)", [str(path)])
        connection.execute("COMMIT")
    report = {
        "source_afac": str(AFAC.relative_to(PATHS.root)).replace("\\", "/"),
        "source_roster": str(ROSTER.relative_to(PATHS.root)).replace("\\", "/"),
        "source_hashes": {"afac": afac_hash, "roster": roster_hash},
        "silver_rows": len(silver), "gold_rows": len(gold),
        "lineage_links": len(bridge),
        "markets": sorted(gold.market_key.tolist()),
        "all_carrier_market_movements_context_only": int(gold.market_movements_all_carriers.sum()),
        "carrier_movements_known": False,
        "agent_eligible": False,
    }
    QUALITY.parent.mkdir(parents=True, exist_ok=True)
    QUALITY.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return report


if __name__ == "__main__":
    print(json.dumps(run(), ensure_ascii=False, indent=2))
